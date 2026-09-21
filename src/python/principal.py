import pandas as pd
import numpy as np
import os
import glob
import ast
import json
import re
import unicodedata
import tifffile as tiff
from docx import Document
from PIL import Image
from openpyxl import load_workbook

from twinelec_paths import assets_dir, blender_scripts_dir, catalogs_dir


HOJA_APOYOS_CRUCETAS = "Apoyos y crucetas"
HOJA_CIMENTACION_MONOBLOQUE = "Cimen. monobloque"
COLUMNAS_APOYOS_CRUCETAS = {
    0: "apoyo_numero",
    1: "referencia_apoyo_origen",
    2: "altura_normalizada_m",
    3: "recrecido_cabeza_m",
    # Se conserva solo para auditar el Excel. La altura del modelo procede de
    # Altura sobre terreno > Cogolla en la hoja de cimentación monobloque.
    4: "altura_total_apoyos_crucetas_m",
    5: "armado_base_origen",
    6: "a_cruceta_m",
    7: "referencia_armado_origen",
    8: "separacion_vertical_crucetas_m",
    9: "altura_cupula_m",
    10: "referencia_cruceta_origen",
    11: "tipo_cruceta_origen",
}

COLUMNAS_CIMENTACION_MONOBLOQUE = {
    0: "apoyo_numero",
    3: "altura_cogolla_m",
}

RUTA_CATALOGO_CONDUCTORES = str(
    catalogs_dir() / "conductors" / "catalogo_conductores.json"
)


def _texto_limpio(valor):
    if valor is None or pd.isna(valor):
        return None
    texto = str(valor).replace("_x000D_", " ").replace("\n", " ").strip()
    texto = re.sub(r"\s+", " ", texto)
    if not texto or all(caracter in "-—– " for caracter in texto):
        return None
    return texto


def _clave_texto(valor):
    texto = _texto_limpio(valor)
    if texto is None:
        return ""
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", texto.lower()).strip()


def _numero_decimal(valor):
    texto = _texto_limpio(valor)
    if texto is None:
        return None
    texto = texto.replace(" ", "").replace(",", ".")
    coincidencia = re.search(r"[-+]?\d+(?:\.\d+)?", texto)
    if coincidencia is None:
        return None
    return float(coincidencia.group(0))


def _clave_conductor(valor):
    """Iguala LA-56, LA 56 y la designación normativa sin perder sus cifras."""
    texto = _texto_limpio(valor)
    if texto is None:
        return ""
    texto = unicodedata.normalize("NFKD", texto).upper()
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return re.sub(r"[^A-Z0-9]+", "", texto)


def crear_conductor_configurado(archivo_excel, carpeta_assets):
    """Resuelve el conductor de Mediciones contra el catálogo documental."""
    mediciones = pd.read_excel(
        archivo_excel, sheet_name="Mediciones", header=None, dtype=object
    )
    nomenclatura = None
    fila_excel = None
    for indice, fila in mediciones.iterrows():
        if any("tipo de conductor" in _clave_texto(valor) for valor in fila):
            valores = [_texto_limpio(valor) for valor in fila]
            valores = [valor for valor in valores if valor]
            if valores:
                nomenclatura = valores[-1]
                fila_excel = indice + 1
            break
    if nomenclatura is None:
        raise ValueError("No se encontró 'Tipo de conductor' en la hoja Mediciones.")

    with open(RUTA_CATALOGO_CONDUCTORES, encoding="utf-8") as archivo:
        catalogo = json.load(archivo)
    clave_buscada = _clave_conductor(nomenclatura)
    coincidencia = None
    for conductor in catalogo["conductores"]:
        nombres = [conductor["designacion"], *conductor.get("aliases", [])]
        if clave_buscada in {_clave_conductor(nombre) for nombre in nombres}:
            coincidencia = conductor
            break
    if coincidencia is None:
        raise ValueError(
            f"El conductor '{nomenclatura}' no figura en {RUTA_CATALOGO_CONDUCTORES}."
        )

    propiedades_composicion = {
        entrada["composicion"]: entrada
        for entrada in catalogo.get("propiedades_mecanicas_por_composicion", [])
    }
    composicion = coincidencia.get("composicion")
    propiedades = propiedades_composicion.get(composicion, coincidencia)
    if "modulo_elasticidad_daN_mm2" not in propiedades:
        raise ValueError(
            f"No hay modulo de elasticidad para la composicion '{composicion}'."
        )
    if "coef_dilatacion_1_C" not in propiedades:
        raise ValueError(
            f"No hay coeficiente de dilatacion para la composicion '{composicion}'."
        )
    fuente_propiedades = coincidencia.get(
        "fuente_propiedades",
        "Tema 1 - Conductores - 2015.pdf, tabla 5, pagina 11",
    )

    documento = {
        "version_esquema": 3,
        "nomenclatura_excel": nomenclatura,
        "fila_excel": fila_excel,
        "designacion": coincidencia["designacion"],
        "composicion": composicion,
        "diametro_mm": coincidencia["diametro_mm"],
        "diametro_m": coincidencia["diametro_mm"] / 1000.0,
        "seccion_mm2": coincidencia["seccion_mm2"],
        "masa_kg_km": coincidencia["masa_kg_km"],
        "peso_daN_m": coincidencia["masa_kg_km"] * 0.98 / 1000.0,
        "modulo_elasticidad_daN_mm2": propiedades[
            "modulo_elasticidad_daN_mm2"
        ],
        "coef_dilatacion_1_C": propiedades["coef_dilatacion_1_C"],
        "propiedades_mecanicas_verificadas": True,
        "fuente_propiedades_mecanicas": fuente_propiedades,
        "numero_fases": 3,
        "fuente_catalogo": "data/catalogs/conductors/catalogo_conductores.json",
    }

    # --- ENRIQUECER CON RESISTENCIAS ELÉCTRICAS (necesarias para el DLR) ---
    # Se buscan R(Tlow) y R(Thigh) en la tabla de resistencias extraída del
    # Excel de EnerFlux (TFG de Manuel). Es GENÉRICO: funciona para cualquier
    # conductor del catálogo que esté en esa tabla.
    try:
        ruta_resistencias = str(
            catalogs_dir() / "conductors" / "resistencias_electricas.json"
        )
        if os.path.exists(ruta_resistencias):
            with open(ruta_resistencias, encoding="utf-8") as archivo:
                tabla_resistencias = json.load(archivo).get("conductores", {})
            clave_nom = _clave_conductor(nomenclatura)
            registro = tabla_resistencias.get(clave_nom)
            if registro is None:
                # Buscar por la designación normativa (p.ej. 47-AL1/8-ST1A)
                clave_designacion = _clave_conductor(coincidencia["designacion"])
                for registro_candidato in tabla_resistencias.values():
                    if _clave_conductor(str(registro_candidato.get("codigo", ""))) == clave_designacion:
                        registro = registro_candidato
                        break
            if registro:
                documento["r_low_ohm_km"] = registro["r_low_ohm_km"]
                documento["r_high_ohm_km"] = registro["r_high_ohm_km"]
                documento["t_low_C"] = registro.get("t_low_C", 25.0)
                documento["t_high_C"] = registro.get("t_high_C", 70.0)
                documento["fuente_resistencias"] = (
                    "data/catalogs/conductors/resistencias_electricas.json"
                )
            else:
                documento["r_low_ohm_km"] = None
                documento["r_high_ohm_km"] = None
                documento["t_low_C"] = 25.0
                documento["t_high_C"] = 70.0
                print(
                    "⚠️ Sin resistencias eléctricas en la tabla para "
                    f"'{nomenclatura}'. El DLR no podrá calcular R(Tavg)."
                )
    except Exception as error:
        print(f"⚠️ No se pudo enriquecer el conductor con resistencias: {error}")

    ruta_salida = os.path.join(carpeta_assets, "conductor_configurado.json")
    with open(ruta_salida, "w", encoding="utf-8") as archivo:
        json.dump(documento, archivo, ensure_ascii=False, indent=2)
        archivo.write("\n")
    return documento


def _clasificar_catalogo(referencia_apoyo):
    clave = _clave_texto(referencia_apoyo)
    if "13000" in clave:
        return "13000"
    if "serie c" in clave:
        return "SERIE_C"
    if "presilla" in clave:
        return "PRESILLA"
    if (
        "andel" in clave
        or re.search(r"\bm\s*(50|60)\b", clave)
        or re.search(r"\bserie\s+g\b", clave)
    ):
        return "ANDEL"
    return None


def _resolver_modelo_armado(referencia_apoyo):
    """Convierte la designación de la columna B en un ID estable para Unity."""
    clave = _clave_texto(referencia_apoyo)
    if not clave:
        return None

    # Conservamos los IDs ya usados para la Serie C actual.
    if "serie c" in clave:
        nivel = re.search(r"\bc\s*(500|1000|2000|3000|4500|7000|9000)\b", clave)
        if nivel:
            return f"SERIE_C_C{nivel.group(1)}"

    # El resto de familias no se codifica mediante una lista cerrada: la
    # designación completa de la columna B se convierte en una clave estable.
    # Así pueden incorporarse nuevos armados simplemente registrando en Unity
    # un prefab con el mismo modeloId, sin modificar este código.
    identificador = re.sub(r"[^A-Z0-9]+", "_", clave.upper()).strip("_")
    return f"ARMADO_{identificador}" if identificador else None


def _resolver_recurso_unity_armado(referencia_apoyo):
    """Modelo visible compartido por las variantes resistentes de cada cabeza."""
    clave = _clave_texto(referencia_apoyo)
    if "serie c" in clave:
        return "ModelosArmado/Cabeza_C"
    if "presilla" in clave and "trenz" in clave:
        return "ModelosArmado/Cabeza_PRESILLA_400_1250_TRENZADOS"
    if "presilla" in clave and re.search(r"\b250\b", clave):
        return "ModelosArmado/Cabeza_PRESILLA_250"
    if "presilla" in clave:
        return "ModelosArmado/Cabeza_PRESILLA_400_1250"
    if re.search(r"\bm\s*60\b", clave):
        return "ModelosArmado/Cabeza_M60"
    if re.search(r"\bm\s*50\b", clave):
        return "ModelosArmado/Cabeza_M50"
    if re.search(r"\bg\b", clave):
        return "ModelosArmado/Cabeza_G"
    return None


def _resolver_cuerpo_modular_armado(referencia_apoyo):
    """Devuelve cabeza y celda inferior para construir la altura total."""
    clave = _clave_texto(referencia_apoyo)
    if "serie c" in clave:
        familia, altura_cabeza, altura_modulo, primer_patron = (
            "C", 4.200, 0.600, "B"
        )
    elif "presilla" in clave and "trenz" in clave:
        familia, altura_cabeza, altura_modulo, primer_patron = (
            "PRESILLA_400_1250_TRENZADOS",
            4.100,
            0.470,
            "B",
        )
    elif "presilla" in clave and re.search(r"\b250\b", clave):
        familia, altura_cabeza, altura_modulo, primer_patron = (
            "PRESILLA_250", 4.000, 0.340, "A"
        )
    elif "presilla" in clave:
        familia, altura_cabeza, altura_modulo, primer_patron = (
            "PRESILLA_400_1250",
            4.100,
            0.470,
            "A",
        )
    elif re.search(r"\bm\s*60\b", clave):
        familia, altura_cabeza, altura_modulo, primer_patron = (
            "M60", 4.000, 0.400, "A"
        )
    elif re.search(r"\bm\s*50\b", clave):
        familia, altura_cabeza, altura_modulo, primer_patron = (
            "M50", 4.000, 0.400, "B"
        )
    elif re.search(r"\bg\b", clave):
        familia, altura_cabeza, altura_modulo, primer_patron = (
            "G", 5.600, 0.700, "A"
        )
    else:
        return None

    ruta = f"ModelosArmado/Modulo_{familia}"
    return {
        "familia": familia,
        "altura_cabeza_m": altura_cabeza,
        "altura_modulo_m": altura_modulo,
        "recurso_modulo_a": f"{ruta}_A",
        "recurso_modulo_b": f"{ruta}_B",
        "primer_patron_bajo_cabeza": primer_patron,
    }


def _normalizar_montaje(armado_base, tipo_cruceta, *referencias_extra):
    clave = " ".join(
        _clave_texto(valor)
        for valor in (armado_base, tipo_cruceta, *referencias_extra)
    )
    if "paso" in clave and "herraje" in clave and "puente" in clave:
        return "PASO_HERRAJE_PUENTE"
    if ("paso" in clave and "lateral" in clave) or "amcx" in clave:
        return "PASO_LATERAL"
    if ("triangulo" in clave and "tg" in clave) or re.search(r"\btg\s*\d+", clave):
        return "TRIANGULO_TG"
    reglas = (
        (("boveda", "capa"), "BOVEDA_CAPA"),
        (("horizontal",), "HORIZONTAL"),
        (("pico",), "BOVEDA_PICO"),
        (("triangulo",), "BOVEDA_TRIANGULO"),
        (("doble", "circuito"), "DOBLE_CIRCUITO"),
        (("tresbolillo",), "TRESBOLILLO"),
        (("bandera",), "BANDERA"),
    )
    for palabras, montaje in reglas:
        if all(palabra in clave for palabra in palabras):
            return montaje
    return "NO_CLASIFICADO"


def buscar_excel_con_apoyos_y_crucetas(carpeta="."):
    candidatos = sorted(
        ruta
        for ruta in glob.glob(os.path.join(carpeta, "*.xlsx"))
        if not os.path.basename(ruta).startswith("~$")
    )
    if not candidatos:
        raise FileNotFoundError("No se encontró ningún archivo .xlsx en la carpeta.")
    # Preferir el Excel ORIGINAL (plantilla) si la copia editada también está:
    # esta función lee la configuración de la línea (apoyos, tendido, T_max);
    # la "_editada" es la SALIDA de las ediciones y no debe usarse de entrada.
    originales = [
        r for r in candidatos
        if not os.path.basename(r).endswith("_editada.xlsx")
    ]
    if len(originales) == 1:
        return originales[0]
    if len(candidatos) > 1:
        raise RuntimeError(
            "Debe haber un único archivo .xlsx en la carpeta. Encontrados: "
            + ", ".join(candidatos)
        )
    return candidatos[0]


def _resolver_hoja_cimentacion_monobloque(archivo_excel):
    """Admite el nombre completo y la abreviatura usada por Andelec."""
    nombres = pd.ExcelFile(archivo_excel).sheet_names
    coincidencias = [
        nombre
        for nombre in nombres
        if "cimen" in _clave_texto(nombre) and "monobloque" in _clave_texto(nombre)
    ]
    if len(coincidencias) != 1:
        raise ValueError(
            "Debe existir una única pestaña de cimentación monobloque. "
            f"Encontradas: {coincidencias or 'ninguna'}"
        )
    return coincidencias[0]


def auditar_apoyos_y_crucetas(archivo_excel=None, ruta_salida=None):
    archivo_excel = archivo_excel or buscar_excel_con_apoyos_y_crucetas()
    tabla = pd.read_excel(
        archivo_excel,
        sheet_name=HOJA_APOYOS_CRUCETAS,
        header=None,
        dtype=object,
    )

    hoja_cimentacion = _resolver_hoja_cimentacion_monobloque(archivo_excel)
    tabla_cimentacion = pd.read_excel(
        archivo_excel,
        sheet_name=hoja_cimentacion,
        header=None,
        dtype=object,
    )

    registros = []
    incidencias = []
    filas_por_apoyo = {}
    alturas_cogolla = {}
    filas_cimentacion_por_apoyo = {}

    for indice_fila, fila in tabla_cimentacion.iterrows():
        valores = {
            nombre: fila.iloc[indice] if indice < len(fila) else None
            for indice, nombre in COLUMNAS_CIMENTACION_MONOBLOQUE.items()
        }
        apoyo_numero = _numero_decimal(valores["apoyo_numero"])
        if apoyo_numero is None or not apoyo_numero.is_integer() or apoyo_numero <= 0:
            continue
        apoyo_numero = int(apoyo_numero)
        fila_excel = indice_fila + 1
        filas_cimentacion_por_apoyo.setdefault(apoyo_numero, []).append(fila_excel)
        altura_cogolla = _numero_decimal(valores["altura_cogolla_m"])
        if altura_cogolla is not None:
            alturas_cogolla[apoyo_numero] = altura_cogolla

    for apoyo_numero, filas in filas_cimentacion_por_apoyo.items():
        if len(filas) > 1:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "hoja": hoja_cimentacion,
                    "fila_excel": filas,
                    "apoyo_numero": apoyo_numero,
                    "campo": "altura_cogolla_m",
                    "mensaje": (
                        f"El apoyo {apoyo_numero} aparece más de una vez en "
                        "cimentación monobloque."
                    ),
                }
            )

    for indice_fila, fila in tabla.iterrows():
        apoyo_numero = _numero_decimal(fila.iloc[0] if len(fila) > 0 else None)
        if apoyo_numero is None:
            continue
        if not apoyo_numero.is_integer() or apoyo_numero <= 0:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "fila_excel": indice_fila + 1,
                    "campo": "apoyo_numero",
                    "mensaje": f"Número de apoyo no válido: {apoyo_numero}",
                }
            )
            continue

        apoyo_numero = int(apoyo_numero)
        valores = {
            nombre: fila.iloc[indice] if indice < len(fila) else None
            for indice, nombre in COLUMNAS_APOYOS_CRUCETAS.items()
        }
        referencia_apoyo = _texto_limpio(valores["referencia_apoyo_origen"])
        catalogo = _clasificar_catalogo(referencia_apoyo)
        armado_base = _texto_limpio(valores["armado_base_origen"])
        tipo_cruceta = _texto_limpio(valores["tipo_cruceta_origen"])
        cuerpo_modular = _resolver_cuerpo_modular_armado(referencia_apoyo)
        tipo_montaje = _normalizar_montaje(
            armado_base,
            tipo_cruceta,
            valores["referencia_armado_origen"],
            valores["referencia_cruceta_origen"],
        )

        registro = {
            "apoyo_numero": apoyo_numero,
            "catalogo": catalogo,
            "referencia_apoyo_origen": referencia_apoyo,
            "modelo_armado_id": _resolver_modelo_armado(referencia_apoyo),
            "recurso_unity_armado": _resolver_recurso_unity_armado(
                referencia_apoyo
            ),
            "cuerpo_modular_armado": cuerpo_modular,
            "altura_normalizada_m": _numero_decimal(
                valores["altura_normalizada_m"]
            ),
            "recrecido_cabeza_m": _numero_decimal(valores["recrecido_cabeza_m"]),
            # Altura real desde el suelo hasta el punto más alto del apoyo.
            # La columna E de "Apoyos y crucetas" queda expresamente descartada.
            "altura_total_m": alturas_cogolla.get(apoyo_numero),
            "armado_base_origen": armado_base,
            "tipo_montaje": tipo_montaje,
            "a_cruceta_m": _numero_decimal(valores["a_cruceta_m"]),
            "referencia_armado_origen": _texto_limpio(
                valores["referencia_armado_origen"]
            ),
            "separacion_vertical_crucetas_m": _numero_decimal(
                valores["separacion_vertical_crucetas_m"]
            ),
            "altura_cupula_m": _numero_decimal(valores["altura_cupula_m"]),
            "referencia_cruceta_origen": _texto_limpio(
                valores["referencia_cruceta_origen"]
            ),
            "tipo_cruceta_origen": tipo_cruceta,
            "es_atirantada": bool(
                (tipo_cruceta and tipo_cruceta.upper().startswith("ATC"))
                or tipo_montaje == "PASO_HERRAJE_PUENTE"
            ),
            "fila_excel": indice_fila + 1,
        }
        registros.append(registro)
        filas_por_apoyo.setdefault(apoyo_numero, []).append(indice_fila + 1)

        if registro["altura_total_m"] is None:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "hoja": hoja_cimentacion,
                    "fila_excel": filas_cimentacion_por_apoyo.get(apoyo_numero),
                    "apoyo_numero": apoyo_numero,
                    "campo": "altura_cogolla_m",
                    "mensaje": (
                        "Falta la altura de cogolla en la pestaña de cimentación "
                        "monobloque."
                    ),
                }
            )

        if referencia_apoyo is None:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "fila_excel": indice_fila + 1,
                    "campo": "referencia_apoyo_origen",
                    "mensaje": "Falta la referencia del apoyo en la columna B.",
                }
            )
        elif catalogo is None:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "fila_excel": indice_fila + 1,
                    "campo": "catalogo",
                    "mensaje": "No se reconoce ANDEL, Serie C o 13000 en la columna B.",
                }
            )
        elif registro["modelo_armado_id"] is None:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "fila_excel": indice_fila + 1,
                    "apoyo_numero": apoyo_numero,
                    "campo": "referencia_apoyo_origen",
                    "mensaje": "No se pudo resolver el modelo de armado de la columna B.",
                }
            )

        if registro["es_atirantada"] and registro["a_cruceta_m"] is None:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "fila_excel": indice_fila + 1,
                    "campo": "a_cruceta_m",
                    "mensaje": "Una cruceta ATC necesita la medida a en la columna G.",
                }
            )
        if registro["tipo_montaje"] == "NO_CLASIFICADO":
            incidencias.append(
                {
                    "nivel": "AVISO",
                    "fila_excel": indice_fila + 1,
                    "campo": "armado_base_origen",
                    "mensaje": "Tipo de montaje todavía no clasificado.",
                }
            )

    for apoyo_numero, filas in filas_por_apoyo.items():
        if len(filas) > 1:
            incidencias.append(
                {
                    "nivel": "ERROR",
                    "fila_excel": filas,
                    "campo": "apoyo_numero",
                    "mensaje": f"El apoyo {apoyo_numero} aparece más de una vez.",
                }
            )

    numeros = sorted(filas_por_apoyo)
    if numeros and numeros != list(range(numeros[0], numeros[-1] + 1)):
        incidencias.append(
            {
                "nivel": "AVISO",
                "fila_excel": None,
                "campo": "apoyo_numero",
                "mensaje": "La numeración de apoyos contiene saltos.",
            }
        )

    resumen_catalogos = {}
    resumen_montajes = {}
    for registro in registros:
        catalogo = registro["catalogo"] or "NO_RECONOCIDO"
        montaje = registro["tipo_montaje"]
        resumen_catalogos[catalogo] = resumen_catalogos.get(catalogo, 0) + 1
        resumen_montajes[montaje] = resumen_montajes.get(montaje, 0) + 1

    resultado = {
        "archivo_excel": os.path.basename(archivo_excel),
        "hoja": HOJA_APOYOS_CRUCETAS,
        "hoja_altura_total": hoja_cimentacion,
        "criterios": {
            "columna_A": "número de apoyo",
            "columna_B": "catálogo y modelo de armado visible en Unity",
            "recrecido": "auditado pero todavía no aplicado al montaje",
            "longitud_cruceta": "medida a para seleccionar el modelo",
            "altura_total": (
                "Altura sobre terreno > Cogolla de la pestaña de cimentación "
                "monobloque; distancia del suelo al punto más alto del apoyo"
            ),
            "separacion_crucetas": "separación vertical",
        },
        "resumen": {
            "numero_registros": len(registros),
            "por_catalogo": resumen_catalogos,
            "por_tipo_montaje": resumen_montajes,
            "errores": sum(i["nivel"] == "ERROR" for i in incidencias),
            "avisos": sum(i["nivel"] == "AVISO" for i in incidencias),
        },
        "registros": registros,
        "incidencias": incidencias,
    }

    if ruta_salida:
        with open(ruta_salida, "w", encoding="utf-8") as archivo:
            json.dump(resultado, archivo, ensure_ascii=False, indent=2)

    return resultado


def cargar_modelos_catalogados(
    ruta_script=None,
):
    ruta_script = ruta_script or str(
        blender_scripts_dir() / "crear_semicrucetas_andel.py"
    )
    with open(ruta_script, "r", encoding="utf-8-sig") as archivo:
        arbol = ast.parse(archivo.read(), filename=ruta_script)

    nombres_tablas = {
        "ALTURAS_ACOPLE",
        "SEMICRUCETAS_ASC",
        "CRUCETAS_ARCX",
        "CRUCETAS_AMC",
        "GANCHO_HERRAJE_PUENTE",
        "MONTAJES_TRIANGULO_TG",
        "BOVEDAS_CAPA",
        "BOVEDAS_PICO",
        "BOVEDAS_TRIANGULO",
        "SEMICRUCETAS_ATIRANTADAS",
        "SERIE_C_SEMICRUCETAS",
        "SERIE_C_CRUCETAS_RECTAS",
        "SERIE_C_ATIRANTADAS",
        "SERIE_C_BOVEDAS_TRIANGULO",
        "SERIE_C_BOVEDAS_CAPA",
        "SERIE_C_BOVEDAS_PICO",
        "CATALOGO_13000_CRUCETAS",
        "PRESILLA_TRESBOLILLO",
        "PRESILLA_RECTAS_MONTAJE_0",
        "PRESILLA_RECTAS_ARRIOSTRADAS",
        "PRESILLA_BOVEDAS_PICO",
    }
    tablas = {}
    for nodo in arbol.body:
        if not isinstance(nodo, ast.Assign) or len(nodo.targets) != 1:
            continue
        destino = nodo.targets[0]
        if isinstance(destino, ast.Name) and destino.id in nombres_tablas:
            tablas[destino.id] = ast.literal_eval(nodo.value)

    ausentes = nombres_tablas - tablas.keys()
    if ausentes:
        raise RuntimeError(
            "No se pudieron leer las tablas del catálogo 3D: "
            + ", ".join(sorted(ausentes))
        )

    alturas = [sufijo for sufijo, _, _ in tablas["ALTURAS_ACOPLE"]]
    modelos = set(tablas["SEMICRUCETAS_ASC"])
    modelos.update(tablas["CRUCETAS_ARCX"])
    modelos.update(tablas["CRUCETAS_AMC"])
    modelos.add(tablas["GANCHO_HERRAJE_PUENTE"])
    modelos.update(tablas["MONTAJES_TRIANGULO_TG"])
    for nombre_tabla in (
        "BOVEDAS_CAPA",
        "BOVEDAS_PICO",
        "BOVEDAS_TRIANGULO",
        "SEMICRUCETAS_ATIRANTADAS",
    ):
        for referencia in tablas[nombre_tabla]:
            modelos.update(f"{referencia}_{altura}" for altura in alturas)

    for nombre_tabla in (
        "SERIE_C_SEMICRUCETAS",
        "SERIE_C_CRUCETAS_RECTAS",
        "SERIE_C_ATIRANTADAS",
        "SERIE_C_BOVEDAS_TRIANGULO",
        "SERIE_C_BOVEDAS_CAPA",
        "SERIE_C_BOVEDAS_PICO",
    ):
        modelos.update(f"C_{referencia}" for referencia in tablas[nombre_tabla])

    modelos.update(
        f"CAT13000_{referencia}"
        for referencia in tablas["CATALOGO_13000_CRUCETAS"]
    )
    modelos.add(tablas["PRESILLA_TRESBOLILLO"])
    modelos.update(tablas["PRESILLA_RECTAS_MONTAJE_0"])
    modelos.update(tablas["PRESILLA_RECTAS_ARRIOSTRADAS"])
    modelos.update(
        referencia for referencia, _, _ in tablas["PRESILLA_BOVEDAS_PICO"]
    )
    return modelos


def _identificador_catalogo(catalogo):
    return {
        "ANDEL": "CATALOGO_ANDEL",
        "SERIE_C": "CATALOGO_SERIE_C",
        "13000": "CATALOGO_13000",
        "PRESILLA": "CATALOGO_PRESILLA",
    }.get(catalogo)


def _referencia_compacta(valor):
    texto = _texto_limpio(valor)
    if texto is None:
        return None
    return re.sub(r"\s+", "", texto.upper())


def _resolver_modelo_cruceta(registro, modelos_catalogados):
    catalogo = registro["catalogo"]
    referencias_origen = [
        registro.get("tipo_cruceta_origen"),
        registro.get("referencia_cruceta_origen"),
        registro.get("referencia_armado_origen"),
    ]
    referencias = [
        referencia
        for referencia in map(_referencia_compacta, referencias_origen)
        if referencia
    ]

    prefijos = {
        "SERIE_C": "C_",
        "13000": "CAT13000_",
        "ANDEL": "",
        "PRESILLA": "",
    }
    prefijo = prefijos.get(catalogo)
    if prefijo is None:
        return None

    for referencia in referencias:
        candidato = f"{prefijo}{referencia}"
        if candidato in modelos_catalogados:
            if catalogo == "ANDEL":
                montaje = registro["tipo_montaje"]
                if registro.get("es_atirantada"):
                    if not candidato.startswith("ATC"):
                        continue
                elif montaje == "PASO_LATERAL":
                    if not candidato.startswith("AMCX-"):
                        continue
                elif montaje == "TRIANGULO_TG":
                    if not candidato.startswith("TG"):
                        continue
                elif montaje in ("DOBLE_CIRCUITO", "HORIZONTAL"):
                    if not candidato.startswith("ARCX-"):
                        continue
                elif montaje in ("TRESBOLILLO", "BANDERA"):
                    if not candidato.startswith("ASC-"):
                        continue
            return candidato

    if catalogo == "ANDEL":
        tipo_montaje = registro["tipo_montaje"]
        a_m = registro.get("a_cruceta_m")
        codigo_por_a = {
            1.250: "12",
            1.500: "15",
            1.750: "17",
            2.000: "20",
        }
        codigo = next(
            (
                valor
                for medida, valor in codigo_por_a.items()
                if a_m is not None and abs(a_m - medida) < 1e-6
            ),
            None,
        )

        # Las bovedas y las ATC se guardan en Blender con una variante H.
        # En el JSON conservamos la referencia base: Unity resolvera H375,
        # H500, H600 o H750 segun el panel local real de M50/M60.
        texto_referencias = " ".join(referencias)
        if tipo_montaje == "TRIANGULO_TG":
            referencias_tg = {
                "TG170": "TG120-AXC",
                "TG120": "TG120-AXC",
                "TG250": "TG250-AXC",
                "TG300": "TG300-AXC",
                "TG400": "TG400-AXC",
            }
            for referencia_tabla, referencia_modelo in referencias_tg.items():
                patron_modelo = referencia_modelo.replace("X", "[X1-4]")
                if (
                    referencia_tabla in texto_referencias
                    or re.search(patron_modelo, texto_referencias)
                ) and referencia_modelo in modelos_catalogados:
                    return referencia_modelo

        if tipo_montaje.startswith("BOVEDA_"):
            coincidencia = re.search(r"B[HFT]\d+-AN[A-Z]", texto_referencias)
            if coincidencia:
                referencia_base = coincidencia.group(0)
                if any(
                    modelo.startswith(f"{referencia_base}_H")
                    for modelo in modelos_catalogados
                ):
                    return referencia_base

        if registro.get("es_atirantada") and codigo:
            coincidencia_serie = re.search(r"(?:ATC|TB)\s*-?(45|90)", texto_referencias)
            serie = coincidencia_serie.group(1) if coincidencia_serie else "45"
            referencia_base = f"ATC{serie}-{codigo}"
            if any(
                modelo.startswith(f"{referencia_base}_H")
                for modelo in modelos_catalogados
            ):
                return referencia_base

        if codigo:
            if tipo_montaje == "PASO_LATERAL":
                familia = "AMCX"
            elif tipo_montaje in ("DOBLE_CIRCUITO", "HORIZONTAL"):
                familia = "ARCX"
            else:
                familia = "ASC"
            candidato = f"{familia}-{codigo}"
            if candidato in modelos_catalogados:
                return candidato

    if catalogo == "SERIE_C" and registro.get("a_cruceta_m") is not None:
        d_m = registro["a_cruceta_m"]
        familias_boveda = {
            "BOVEDA_TRIANGULO": "C_BT",
            "BOVEDA_CAPA": "C_BH",
            "HORIZONTAL": "C_BH",
            "BOVEDA_PICO": "C_BF",
        }
        prefijo_boveda = familias_boveda.get(registro["tipo_montaje"])
        if prefijo_boveda:
            candidatos = sorted(
                modelo
                for modelo in modelos_catalogados
                if modelo.startswith(prefijo_boveda)
            )
            coincidencias = []
            for candidato in candidatos:
                numeros = re.findall(r"\d+", candidato)
                if numeros and abs(float(numeros[0]) / 10 - d_m) < 1e-9:
                    coincidencias.append(candidato)
            if len(coincidencias) == 1:
                return coincidencias[0]

    return None


def _milimetros(valor):
    return None if valor is None else int(round(float(valor) * 1000))


def _crear_configuracion_id(registro, modelo_cruceta):
    partes = [
        "CFG1",
        registro["catalogo"] or "CATALOGO_DESCONOCIDO",
        registro.get("modelo_armado_id") or "ARMADO_NO_RESUELTO",
        registro["tipo_montaje"],
        modelo_cruceta or "MODELO_NO_RESUELTO",
    ]
    dimensiones = (
        ("A", registro.get("a_cruceta_m")),
        ("SEP", registro.get("separacion_vertical_crucetas_m")),
        ("H", registro.get("altura_total_m")),
        ("REC", registro.get("recrecido_cabeza_m")),
        ("CUP", registro.get("altura_cupula_m")),
    )
    for etiqueta, valor in dimensiones:
        valor_mm = _milimetros(valor)
        if valor_mm is not None:
            partes.append(f"{etiqueta}{valor_mm}")
    identificador = "__".join(partes).upper()
    return re.sub(r"[^A-Z0-9_-]+", "_", identificador)


def crear_apoyos_configurados(
    auditoria,
    ruta_salida,
    ruta_catalogo=None,
):
    ruta_catalogo = ruta_catalogo or str(
        blender_scripts_dir() / "crear_semicrucetas_andel.py"
    )
    modelos_catalogados = cargar_modelos_catalogados(ruta_catalogo)
    apoyos = []
    errores = []

    for registro in auditoria["registros"]:
        catalogo_id = _identificador_catalogo(registro["catalogo"])
        modelo_cruceta = _resolver_modelo_cruceta(registro, modelos_catalogados)
        errores_apoyo = []

        if catalogo_id is None:
            errores_apoyo.append("Catálogo no reconocido.")
        if registro.get("modelo_armado_id") is None:
            errores_apoyo.append("Modelo de armado de la columna B no reconocido.")
        if registro.get("recurso_unity_armado") is None:
            errores_apoyo.append("No existe un recurso Unity para ese tipo de armado.")
        if registro.get("cuerpo_modular_armado") is None:
            errores_apoyo.append("No se pudo resolver el cuerpo modular del armado.")
        if registro["tipo_montaje"] == "NO_CLASIFICADO":
            errores_apoyo.append("Tipo de montaje no reconocido.")
        if modelo_cruceta is None:
            errores_apoyo.append(
                "No se encontró un modelo 3D inequívoco para la cruceta indicada."
            )

        apoyo_id = f"APOYO_{registro['apoyo_numero']:04d}"
        configuracion_id = _crear_configuracion_id(registro, modelo_cruceta)
        apoyo = {
            "apoyo_id": apoyo_id,
            "configuracion_id": configuracion_id,
            "catalogo_id": catalogo_id,
            "armado": {
                "modelo_id": registro["modelo_armado_id"],
                "referencia_origen": registro["referencia_apoyo_origen"],
                "recurso_unity": registro["recurso_unity_armado"],
                **(registro["cuerpo_modular_armado"] or {}),
            },
            "apoyo": {
                "numero": registro["apoyo_numero"],
                "altura_normalizada_m": registro["altura_normalizada_m"],
                "altura_total_m": registro["altura_total_m"],
                "recrecido_cabeza_m": registro["recrecido_cabeza_m"],
            },
            "montaje": {
                "montaje_id": f"MONTAJE_{registro['tipo_montaje']}",
                "tipo": registro["tipo_montaje"],
                "referencia_armado": registro["referencia_armado_origen"],
                "separacion_vertical_m": registro[
                    "separacion_vertical_crucetas_m"
                ],
                "altura_cupula_m": registro["altura_cupula_m"],
                # La hoja actual no contiene una columna de orientacion.
                "lado_inicial": 1,
            },
            "cruceta": {
                "modelo_id": modelo_cruceta,
                "tipo_origen": registro["tipo_cruceta_origen"],
                "referencia_origen": registro["referencia_cruceta_origen"],
                "a_m": registro["a_cruceta_m"],
                "es_atirantada": registro["es_atirantada"],
            },
            "origen": {
                "fila_excel": registro["fila_excel"],
                "apoyo_numero": registro["apoyo_numero"],
                "archivo_auditoria": "auditoria_apoyos_crucetas.json",
            },
            "validacion": {
                "estado": "VALIDO" if not errores_apoyo else "ERROR",
                "errores": errores_apoyo,
            },
        }
        apoyos.append(apoyo)
        errores.extend(
            {"apoyo_id": apoyo_id, "mensaje": mensaje}
            for mensaje in errores_apoyo
        )

    documento = {
        "version_esquema": 1,
        "fuente": os.path.basename(auditoria["archivo_excel"]),
        "catalogo_modelos": "blender/scripts/crear_semicrucetas_andel.py",
        "resumen": {
            "numero_apoyos": len(apoyos),
            "configuraciones_unicas": len(
                {apoyo["configuracion_id"] for apoyo in apoyos}
            ),
            "apoyos_validos": sum(
                apoyo["validacion"]["estado"] == "VALIDO" for apoyo in apoyos
            ),
            "apoyos_con_error": sum(
                apoyo["validacion"]["estado"] == "ERROR" for apoyo in apoyos
            ),
        },
        "apoyos": apoyos,
        "errores": errores,
    }

    # Enriquecimiento con tipo de apoyo y cadena de aisladores (Unity):
    # suspension -> cadena vertical; amarre/anclaje -> cadena tras el cable.
    # Sin esto, el enriquecimiento se perderia cada vez que se regenera el JSON.
    try:
        from cambiar_cadena_conductor import enriquecer_documento_apoyos
        resumen_cadenas = enriquecer_documento_apoyos(
            documento, auditoria["archivo_excel"])
        print("  [principal] Cadenas/tipos sincronizados: %d apoyos, %d cadenas"
              % (resumen_cadenas["tipos_aplicados"],
                 resumen_cadenas["cadenas_asignadas_por_defecto"]))
    except Exception as error:
        print("  ⚠️ No se pudo enriquecer apoyos_configurados.json:", error)

    with open(ruta_salida, "w", encoding="utf-8") as archivo:
        json.dump(documento, archivo, ensure_ascii=False, indent=2)
    return documento

def _es_borde_separador(borde):
    return borde is not None and borde.style in {"medium", "thick", "double"}


def _tension_desde_flecha(flecha_m, a_m, b_m, peso_daN_m):
    """Invierte la ecuacion de flecha del programa Pascal por biseccion."""
    if min(flecha_m, a_m, b_m, peso_daN_m) <= 0:
        return None

    def flecha(tension):
        base = (a_m * b_m * peso_daN_m) / (8.0 * tension)
        correccion = 1.0 + (a_m**2 * peso_daN_m**2) / (48.0 * tension**2)
        return base * correccion

    inferior = 1e-6
    superior = 1.0
    while flecha(superior) > flecha_m and superior < 1e9:
        superior *= 2.0
    if superior >= 1e9:
        return None
    for _ in range(100):
        centro = (inferior + superior) * 0.5
        if flecha(centro) > flecha_m:
            inferior = centro
        else:
            superior = centro
    return (inferior + superior) * 0.5


def _extraer_pares_temperatura(hoja):
    pares = []
    for columna_tension in range(7, hoja.max_column, 2):
        cabecera_tension = _clave_texto(hoja.cell(3, columna_tension).value)
        cabecera_flecha = _clave_texto(hoja.cell(3, columna_tension + 1).value)
        temperatura = _numero_decimal(hoja.cell(2, columna_tension).value)
        if temperatura is None or not cabecera_tension.startswith("t dan"):
            continue
        if not cabecera_flecha.startswith("f m"):
            continue
        pares.append((float(temperatura), columna_tension, columna_tension + 1))
    if not pares:
        raise ValueError("No se detectaron columnas T/F con temperatura.")
    return pares


def procesar_datos_fisicos(carpeta_assets, datos_crudos, conductor=None,
                           archivo=None):
    """Exporta la referencia mecanica; Unity incorpora los anclajes reales.
    'archivo': opcional, ruta del .xlsx concreto (p.ej. el de trabajo del
    gemelo). Si no se pasa, usa el primer *.xlsx del directorio actual."""
    if archivo is not None:
        archivos_excel = [archivo]
    else:
        archivos_excel = [
            ruta for ruta in glob.glob("*.xlsx")
            if not os.path.basename(ruta).startswith("~$")
        ]
    if not archivos_excel:
        return

    archivo = archivos_excel[0]
    try:
        if conductor is None:
            ruta_conductor = os.path.join(carpeta_assets, "conductor_configurado.json")
            with open(ruta_conductor, encoding="utf-8") as entrada:
                conductor = json.load(entrada)

        libro = load_workbook(archivo, data_only=True, read_only=False)
        hoja = libro["T. tend. cond. fase"]
        pares_temperatura = _extraer_pares_temperatura(hoja)
        numero_vanos = len(datos_crudos) - 1
        filas_datos = []
        for fila in range(4, hoja.max_row + 1):
            if _numero_decimal(hoja.cell(fila, 4).value) is None:
                continue
            if any(
                _numero_decimal(hoja.cell(fila, columna_t).value) is not None
                for _, columna_t, _ in pares_temperatura
            ):
                filas_datos.append(fila)
        if len(filas_datos) < numero_vanos:
            raise ValueError(
                f"Hay {len(filas_datos)} filas fisicas para {numero_vanos} vanos."
            )
        filas_datos = filas_datos[:numero_vanos]

        inicios_por_borde = []
        for indice, fila in enumerate(filas_datos):
            if indice > 0:
                fila_anterior = filas_datos[indice - 1]
                separador = any(
                    _es_borde_separador(hoja.cell(fila, columna).border.top)
                    or _es_borde_separador(
                        hoja.cell(fila_anterior, columna).border.bottom
                    )
                    for columna in range(1, 7)
                )
                if separador:
                    inicios_por_borde.append(indice)

        if not inicios_por_borde:
            grupos_indices = [[indice] for indice in range(numero_vanos)]
            criterio_cantones = "sin_separadores_cada_vano_independiente"
        else:
            cortes = [0, *inicios_por_borde, numero_vanos]
            grupos_indices = [
                list(range(cortes[i], cortes[i + 1]))
                for i in range(len(cortes) - 1)
            ]
            criterio_cantones = "separadores_horizontales_gruesos"

        canton_por_vano = {}
        for numero_canton, indices in enumerate(grupos_indices, start=1):
            for indice in indices:
                canton_por_vano[indice] = numero_canton

        peso = float(conductor["peso_daN_m"])
        vanos = []
        avisos = []
        if not conductor.get("propiedades_mecanicas_verificadas", True):
            avisos.append(
                "Las propiedades mecanicas de este conductor requieren "
                "verificacion documental antes de uso de ingenieria."
            )
        for indice in range(numero_vanos):
            p1 = datos_crudos[indice]
            p2 = datos_crudos[indice + 1]
            fila = filas_datos[indice]
            a_utm = float(np.hypot(p2["x"] - p1["x"], p2["y"] - p1["y"]))
            h_cotas = float(p2["cota"] - p1["cota"])
            a_excel = _numero_decimal(hoja.cell(fila, 4).value) or a_utm
            h_excel = _numero_decimal(hoja.cell(fila, 5).value)
            if h_excel is None:
                h_excel = h_cotas
            b_excel = float(np.hypot(a_excel, h_excel))

            estados = []
            for temperatura, columna_t, columna_f in pares_temperatura:
                tension = _numero_decimal(hoja.cell(fila, columna_t).value)
                flecha = _numero_decimal(hoja.cell(fila, columna_f).value)
                if tension is None or tension <= 0 or flecha is None or flecha <= 0:
                    continue
                tension_invertida = _tension_desde_flecha(
                    flecha, a_excel, b_excel, peso
                )
                diferencia_pct = None
                if tension_invertida is not None:
                    diferencia_pct = 100.0 * (tension_invertida - tension) / tension
                estados.append({
                    "temperatura_C": temperatura,
                    "tension_daN": tension,
                    "flecha_m": flecha,
                    "tension_invertida_daN": tension_invertida,
                    "diferencia_inversion_pct": diferencia_pct,
                })

            if not estados:
                raise ValueError(f"El vano {indice + 1} no contiene estados T/F validos.")
            vanos.append({
                "indice": indice + 1,
                "tramo": _texto_limpio(hoja.cell(fila, 1).value) or f"{indice + 1}-{indice + 2}",
                "canton_id": canton_por_vano[indice],
                "fila_excel": fila,
                "a_utm_m": a_utm,
                "h_cotas_m": h_cotas,
                "a_excel_m": a_excel,
                "h_excel_m": h_excel,
                "vano_regulacion_excel_m": _numero_decimal(hoja.cell(fila, 6).value),
                "estados_excel": estados,
            })

        cantones = []
        for numero_canton, indices in enumerate(grupos_indices, start=1):
            temperaturas_comunes = None
            estados_por_vano = []
            for indice in indices:
                por_temperatura = {
                    estado["temperatura_C"]: estado
                    for estado in vanos[indice]["estados_excel"]
                }
                estados_por_vano.append(por_temperatura)
                disponibles = set(por_temperatura)
                temperaturas_comunes = (
                    disponibles
                    if temperaturas_comunes is None
                    else temperaturas_comunes & disponibles
                )
            if not temperaturas_comunes:
                raise ValueError(
                    f"El canton {numero_canton} no tiene una temperatura comun "
                    "con tension y flecha validas en todos sus vanos."
                )

            temperaturas_ordenadas = sorted(temperaturas_comunes)
            centro_intervalo = (
                temperaturas_ordenadas[0] + temperaturas_ordenadas[-1]
            ) * 0.5
            temperatura_referencia = min(
                temperaturas_ordenadas,
                key=lambda temperatura: (
                    abs(temperatura - centro_intervalo), temperatura
                ),
            )
            tensiones = [
                por_temperatura[temperatura_referencia]["tension_daN"]
                for por_temperatura in estados_por_vano
            ]
            tension_referencia = tensiones[0]
            if any(
                abs(tension - tension_referencia) > 0.01
                for tension in tensiones[1:]
            ):
                detalle = ", ".join(f"{tension:g}" for tension in tensiones)
                raise ValueError(
                    f"El canton {numero_canton} no comparte una unica tension a "
                    f"{temperatura_referencia:g} C: [{detalle}] daN."
                )

            for indice, por_temperatura in zip(indices, estados_por_vano):
                referencia = por_temperatura[temperatura_referencia]
                vanos[indice]["temperatura_referencia_C"] = temperatura_referencia
                vanos[indice]["tension_referencia_daN"] = tension_referencia
                vanos[indice]["flecha_referencia_m"] = referencia["flecha_m"]
                diferencia = referencia["diferencia_inversion_pct"]
                if diferencia is not None and abs(diferencia) > 10.0:
                    avisos.append(
                        f"Vano {indice + 1}: T Excel e inversion de F difieren "
                        f"{diferencia:.1f}% a {temperatura_referencia:g} C."
                    )
            cantones.append({
                "id": numero_canton,
                "vanos": [indice + 1 for indice in indices],
                "temperatura_referencia_C": temperatura_referencia,
                "tension_referencia_daN": tension_referencia,
            })

        documento = {
            "version_esquema": 3,
            "fuente_excel": os.path.basename(archivo),
            "hoja": "T. tend. cond. fase",
            "hipotesis": "PESO_PROPIO_SIN_VIENTO_NI_HIELO",
            "criterio_cantones": criterio_cantones,
            "propiedades_conductor": {
                "designacion": conductor["designacion"],
                "composicion": conductor["composicion"],
                "seccion_mm2": conductor["seccion_mm2"],
                "masa_kg_km": conductor["masa_kg_km"],
                "peso_daN_m": conductor["peso_daN_m"],
                "modulo_elasticidad_daN_mm2": conductor[
                    "modulo_elasticidad_daN_mm2"
                ],
                "coef_dilatacion_1_C": conductor["coef_dilatacion_1_C"],
                "propiedades_mecanicas_verificadas": conductor.get(
                    "propiedades_mecanicas_verificadas", True
                ),
                "fuente_propiedades_mecanicas": conductor[
                    "fuente_propiedades_mecanicas"
                ],
            },
            "cantones": cantones,
            "vanos": vanos,
            "avisos": avisos,
        }
        ruta_salida = os.path.join(carpeta_assets, "fisica_linea.json")
        with open(ruta_salida, "w", encoding="utf-8") as salida:
            json.dump(documento, salida, ensure_ascii=False, indent=2)
            salida.write("\n")
        print(
            "Modelo mecanico continuo creado: "
            f"{len(vanos)} vanos, {len(cantones)} cantones, peso propio solamente."
        )
        for aviso in avisos:
            print(f"AVISO mecanico: {aviso}")
    except Exception as error:
        print(f"Error al procesar fisica: {error}")

def extraer_datos_terreno_local(min_x, min_y, max_x, max_y):
    archivos_tif = glob.glob("PNOA_MDT05*.tif")
    if not archivos_tif:
        print("❌ No hay archivos PNOA_MDT05 en la carpeta.")
        return None

    res_ign = 5.0
    ancho_px = int(round((max_x - min_x) / res_ign))
    alto_px = int(round((max_y - min_y) / res_ign))
    
    #aquí empieza a haber ceros
    matriz_mosaico = np.zeros((alto_px, ancho_px), dtype=np.float32)

    for ruta_tif in archivos_tif:
        try:
            with tiff.TiffFile(ruta_tif) as tif:
                tags = tif.pages[0].tags
                origin_x = tags[33922].value[3]
                origin_y = tags[33922].value[4]
                mapa_tif = tif.asarray()
                tif_alto, tif_ancho = mapa_tif.shape
                #lo paso a metros porque tif_alto y tif_ancho están en pixeles
                tif_max_x = origin_x + (tif_ancho * res_ign)
                tif_min_y = origin_y - (tif_alto * res_ign)

                if not (origin_x > max_x or tif_max_x < min_x or origin_y < min_y or tif_min_y > max_y):
                    for r in range(alto_px):
                        #en UTM
                        curr_y = max_y - (r * res_ign)
                        if tif_min_y <= curr_y <= origin_y:
                            #EN PIXEL
                            row_tif = int(round((origin_y - curr_y) / res_ign))
                            #AHORA LO MISMO CON LAS COLUMNAS
                            for c in range(ancho_px):
                                curr_x = min_x + (c * res_ign)
                                if origin_x <= curr_x <= tif_max_x:
                                    col_tif = int(round((curr_x - origin_x) / res_ign))
                                    if 0 <= row_tif < tif_alto and 0 <= col_tif < tif_ancho:
                                        matriz_mosaico[r, c] = mapa_tif[row_tif, col_tif]
        except Exception as e:
            print(f"⚠️ Error leyendo {ruta_tif}: {e}")

    # --- BLOQUE DE DEPURACIÓN: GUARDAR MATRIZ EN CSV ---
    try:
        # Guardamos la matriz en un CSV para que lo abras en Excel
        # Cada celda será una altura. Si ves 0.0, ahí está el barranco.
        if os.environ.get("TWINELEC_DEBUG_EXPORTS") == "1":
            np.savetxt("matriz_debug.csv", matriz_mosaico, delimiter=";", fmt="%.2f")
            print("📝 Exportada matriz_debug.csv (TWINELEC_DEBUG_EXPORTS=1).")
    except Exception as e:
        print(f"⚠️ No se pudo crear el archivo de notas: {e}")
    # ---------------------------------------------------      
    return matriz_mosaico

# --- NUEVA FUNCIÓN PARA LA TEXTURA ---
def extraer_textura_local(min_x, min_y, max_x, max_y):
    archivos_foto = glob.glob("PNOA_MA*.tif")
    if not archivos_foto:
        print("❌ No hay archivos de IMAGEN (PNOA-MA) en la carpeta.")
        return None

    res_foto = 0.25 
    ancho_px = int(round((max_x - min_x) / res_foto))
    alto_px = int(round((max_y - min_y) / res_foto))
    
    print(f"📸 Generando textura de {ancho_px}x{alto_px}...")

    for ruta_tif in archivos_foto:
        # 1) tifffile (rapido; necesita imagecodecs para TIF con compresion JPEG)
        try:
            with tiff.TiffFile(ruta_tif) as tif:
                ox, oy = tif.pages[0].tags[33922].value[3], tif.pages[0].tags[33922].value[4]
                img_tif = tif.asarray()
                if img_tif.shape[2] == 4: img_tif = img_tif[:,:,:3]
                
                h, w, _ = img_tif.shape
                # Si el archivo cubre la zona, recortamos y devolvemos
                if not (ox > max_x or (ox + w*res_foto) < min_x or oy < min_y or (oy - h*res_foto) > max_y):
                    c1, c2 = int(round((min_x-ox)/res_foto)), int(round((max_x-ox)/res_foto))
                    r1, r2 = int(round((oy-max_y)/res_foto)), int(round((oy-min_y)/res_foto))
                    return Image.fromarray(img_tif[max(0,r1):r2, max(0,c1):c2])
        except Exception as e:
            print(f"⚠️ Error en imagen (tifffile): {e}")
        # 2) PIL (respaldo para el .exe: lee TIF con compresion JPEG sin
        #    imagecodecs). El ortofoto PNOA viene comprimido JPEG.
        try:
            # El ortofoto supera el limite anti-bomba de PIL (2.240 Mpx).
            Image.MAX_IMAGE_PIXELS = None
            with Image.open(ruta_tif) as im:
                tie = im.tag_v2.get(33922)
                if not tie or len(tie) < 6:
                    print("⚠️ PIL no encontró la georreferencia del ortofoto.")
                    continue
                ox = float(tie[3]); oy = float(tie[4])
                c1 = max(0, int(round((min_x - ox) / res_foto)))
                c2 = min(im.width, int(round((max_x - ox) / res_foto)))
                r1 = max(0, int(round((oy - max_y) / res_foto)))
                r2 = min(im.height, int(round((oy - min_y) / res_foto)))
                if c2 - c1 <= 0 or r2 - r1 <= 0:
                    continue
                if im.mode in ("CMYK", "RGBA"):
                    im = im.convert("RGB")
                return im.crop((c1, r1, c2, r2)).convert("RGB")
        except Exception as e:
            print(f"⚠️ Error en imagen (PIL): {e}")
    return None

def procesar_word_utm():
    archivos = glob.glob("*UTM*.doc*")
    if not archivos: return
    doc = Document(archivos[0])
    datos_crudos = []
    for i, fila in enumerate(doc.tables[0].rows):
        if i == 0: continue 
        celdas = [c.text.replace(',', '.') for c in fila.cells]
        datos_crudos.append({'id': celdas[0], 'x': float(celdas[1]), 'y': float(celdas[2]), 'cota': float(celdas[3])})

    ux, uy = [d['x'] for d in datos_crudos], [d['y'] for d in datos_crudos]
    margen = 500 
    min_x, max_x = min(ux) - margen, max(ux) + margen
    min_y, max_y = min(uy) - margen, max(uy) + margen
    
    # 1. Definimos la ruta de la carpeta de Unity
    carpeta_assets = str(assets_dir())
    if not os.path.exists(carpeta_assets):
        os.makedirs(carpeta_assets)

    # --- CONDUCTOR DE FASE: NOMENCLATURA Y DIÁMETRO REAL ---
    conductor = None
    try:
        excel_mediciones = buscar_excel_con_apoyos_y_crucetas()
        conductor = crear_conductor_configurado(excel_mediciones, carpeta_assets)
        print(
            "✅ Conductor configurado: "
            f"{conductor['nomenclatura_excel']} -> {conductor['diametro_mm']:.2f} mm, "
            f"{conductor['numero_fases']} fases."
        )
    except Exception as error:
        print(f"❌ No se pudo configurar el conductor: {error}")

    # --- AUDITORÍA DEL ARMADO Y LAS CRUCETAS ---
    try:
        excel_apoyos = buscar_excel_con_apoyos_y_crucetas()
        ruta_auditoria = os.path.join(
            carpeta_assets,
            "auditoria_apoyos_crucetas.json",
        )
        auditoria = auditar_apoyos_y_crucetas(excel_apoyos, ruta_auditoria)
        resumen = auditoria["resumen"]
        print(
            "✅ Auditoría de apoyos completada: "
            f"{resumen['numero_registros']} registros, "
            f"{resumen['errores']} errores y {resumen['avisos']} avisos."
        )
        ruta_configuraciones = os.path.join(
            carpeta_assets,
            "apoyos_configurados.json",
        )
        configuraciones = crear_apoyos_configurados(
            auditoria,
            ruta_configuraciones,
        )
        resumen_configuraciones = configuraciones["resumen"]
        print(
            "✅ Configuraciones de apoyos creadas: "
            f"{resumen_configuraciones['apoyos_validos']} válidas y "
            f"{resumen_configuraciones['apoyos_con_error']} con error."
        )
    except Exception as error:
        print(f"❌ No se pudo auditar la hoja de apoyos y crucetas: {error}")

    # --- PROCESAR FÍSICA ---
    procesar_datos_fisicos(carpeta_assets, datos_crudos, conductor)

    # 2. Relieve
    matriz_alturas = extraer_datos_terreno_local(min_x, min_y, max_x, max_y)
    if matriz_alturas is not None:
        h_min, h_max = matriz_alturas.min(), matriz_alturas.max()
        print(f"hmin{h_min},hmax{h_max}")
        matriz_norm = (((matriz_alturas - h_min) / (h_max - h_min) )* 255).astype(np.uint8)

        # --- BLOQUE DE DEPURACIÓN: GUARDAR MATRIZ NORMALIZADA (PÍXELES) ---
        try:
            # Guardamos la matriz de 0-255 en un CSV
            # Si ves un 0, es negro puro (punto más bajo). Si ves 255, es blanco (punto más alto).
            if os.environ.get("TWINELEC_DEBUG_EXPORTS") == "1":
                np.savetxt("matriz_pixeles_debug.csv", matriz_norm,
                           delimiter=";", fmt="%d")
                print(
                    "📝 Exportada matriz_pixeles_debug.csv "
                    f"(rango: {h_min:.2f} m a {h_max:.2f} m)."
                )
        except Exception as e:
            print(f"⚠️ No se pudo crear el archivo de píxeles: {e}")
        # ------------------------------------------------------------------

        # 1. Definimos la ruta de la carpeta y la CREAMOS antes de guardar
        carpeta_assets = str(assets_dir())
    
        if not os.path.exists(carpeta_assets):
            os.makedirs(carpeta_assets) # Esta línea TIENE que tener más espacios que el if
        
        # 2. Ahora que la carpeta existe seguro, guardamos el relieve
        Image.fromarray(matriz_norm).save(os.path.join(carpeta_assets, "terreno_real.png"))

        # 2.5. Textura (AÑADIDO)
        img_foto = extraer_textura_local(min_x, min_y, max_x, max_y)
        if img_foto:
            img_foto.save(os.path.join(carpeta_assets, "textura_real.jpg"), quality=90)
            print("✅ Textura guardada.")

        # 3. Datos de apoyos y ángulos
        origin_x, origin_z, origin_y = datos_crudos[0]['x'], datos_crudos[0]['y'], datos_crudos[0]['cota']
        datos_finales = []
        puntos_lista = [] 

        # Calculamos coordenadas relativas al primer apoyo
        for d in datos_crudos:
            rel_x, rel_z, rel_y = d['x'] - origin_x, d['y'] - origin_z, d['cota'] - origin_y
            puntos_lista.append((rel_x, rel_z, rel_y))
            datos_finales.append([d['id'], round(rel_x, 3), round(rel_z, 3), round(rel_y, 3)])

        # Función interna para calcular el ángulo de giro de cada torre
        def calcular_angulo_viva(p1, p2, p3):
            v1 = np.array([p1[0] - p2[0], p1[1] - p2[1]])
            v2 = np.array([p3[0] - p2[0], p3[1] - p2[1]])
            dot = np.dot(v1, v2)
            mag1 = np.linalg.norm(v1)
            mag2 = np.linalg.norm(v2)
            # Calculamos el ángulo y lo pasamos a grados
            return round(np.degrees(np.arccos(np.clip(dot / (mag1 * mag2), -1.0, 1.0))), 2)

        # Aplicamos el ángulo a cada fila
        for i in range(len(datos_finales)):
            if i == 0 or i == len(datos_finales) - 1:
                # Torres de principio y fin de línea (miran al frente)
                datos_finales[i].append(180.0)
            else:
                # Torres intermedias con ángulo de deflexión
                angulo = calcular_angulo_viva(puntos_lista[i-1], puntos_lista[i], puntos_lista[i+1])
                datos_finales[i].append(angulo)

        # Guardamos el CSV de apoyos en la carpeta de Assets de Unity
        pd.DataFrame(datos_finales, columns=['Apoyo', 'X', 'Z', 'Y', 'Angulo']).to_csv(
            os.path.join(carpeta_assets, 'coordenadas_linea.csv'), index=False
        )       
        
        pd.DataFrame([{
            'min_x': min_x, 'max_x': max_x, 'min_y': min_y, 'max_y': max_y, 
            'h_min': h_min, 'h_max': h_max,
            'ancho': max_x - min_x, 'largo': max_y - min_y,
            'offset_x': origin_x - min_x,
            'offset_z': origin_z - min_y
        }]).to_csv(os.path.join(carpeta_assets, 'info_terreno.csv'), index=False)

        # --- GENERAR CASO DE REFERENCIA DLR (automático para CUALQUIER línea) ---
        # Al final de la cadena de extracción se genera el caso de referencia
        # con la T_máx detectada del Excel, el conductor enriquecido con sus
        # resistencias y los apoyos UTM reales. Cambias de Excel -> se
        # regenera el caso de ESA línea. Nada queda fijado a un caso concreto.
        try:
            import dlr

            # Detectar huso UTM del terreno (p.ej. PNOA_*_HU30_* -> 30).
            huso = 30  # valor por defecto (España peninsular)
            archivos_huso = glob.glob("PNOA*HU*.tif") + glob.glob("PNOA*HU*.TIF")
            for archivo_huso in archivos_huso:
                coincidencia_huso = re.search(r"HU(\d{2})", archivo_huso)
                if coincidencia_huso:
                    huso = int(coincidencia_huso.group(1))
                    break

            excel_dlr = buscar_excel_con_apoyos_y_crucetas()
            t_max_excel = dlr.detectar_temperatura_maxima_excel(excel_dlr)
            if t_max_excel is None:
                t_max_excel = 85.0
                print(
                    "⚠️ No se pudo detectar T_máx del Excel; "
                    f"fallback = {t_max_excel:.0f} °C"
                )

            # Persistir el límite térmico detectado junto al conductor. De
            # este modo clima.py y Unity usan la misma Ts que el Excel.
            if conductor is not None:
                conductor["t_limite_C"] = float(t_max_excel)
                ruta_conductor_dlr = os.path.join(
                    carpeta_assets, "conductor_configurado.json"
                )
                with open(ruta_conductor_dlr, "w", encoding="utf-8") as archivo:
                    json.dump(conductor, archivo, ensure_ascii=False, indent=2)
                    archivo.write("\n")

            # El conductor enriquecido ya trae r_low_ohm_km y r_high_ohm_km.
            if conductor is not None and conductor.get("r_low_ohm_km"):
                dlr.generar_caso_referencia(
                    conductor=conductor,
                    apoyos_utm=datos_crudos,
                    huso=huso,
                    hemisferio="N",
                    ts_C=t_max_excel,
                    modo="ieee738",
                    ruta_salida=os.path.join(
                        carpeta_assets, "caso_referencia_nivel1b.json"
                    ),
                    nombre_caso=(
                        f"Caso generado por principal.py (Ts={t_max_excel:.0f} °C)"
                    ),
                    descripcion=(
                        "Caso de referencia DLR generado automáticamente para "
                        "la línea del Excel procesado, con su T_máx detectada "
                        "de la hoja de tendido y sus datos (conductor, "
                        "apoyos UTM, geometría)."
                    ),
                )
                print(
                    f"✅ Caso DLR generado: caso_referencia_nivel1b.json "
                    f"(Ts={t_max_excel:.0f} °C)"
                )
            else:
                print(
                    "⚠️ No se generó el caso DLR: el conductor no tiene "
                    "resistencias eléctricas (r_low_ohm_km / r_high_ohm_km)."
                )
        except Exception as error:
            print(f"⚠️ No se pudo generar el caso DLR: {error}")

        print(f"✅ Todo generado con éxito.")

if __name__ == "__main__":
    # Python hereda a veces CP-1252 en Windows; sin UTF-8 los prints con
    # simbolos (✅/❌/⚠️) lanzan UnicodeEncodeError y el script ABORTA antes de
    # generar coordenadas_linea.csv -> la app muestra "Revisa los archivos de
    # entrada". Este envoltorio garantiza UTF-8 aunque no lleguen las variables
    # de entorno PYTHONUTF8/PYTHONIOENCODING.
    import sys
    import io
    try:
        if sys.stdout.encoding and "utf" not in sys.stdout.encoding.lower():
            sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                          errors="replace")
    except Exception:
        pass
    procesar_word_utm()
