# -*- coding: utf-8 -*-
"""
editar_tension.py
=================
Edicion MANUAL de la tension de un tramo en 'T. reg. cond. fase' del gemelo
digital. A partir de ese valor se recalcula TODO en cascada:

  1. Tension EDS del vano de regulacion del tramo (inversa de la ecuacion de
     cambio de estado, Subprograma 6 de Pascal): la T nueva define el estado
     de referencia (t0, T0, p0) del grupo.
  2. RESTO DE TENSIONES del grupo (todas las hipotesis: TVF, THF, TVM,
     15C+V, 0C+H, 50C y V+H si el Excel la trae) -> 'T. reg. cond. fase'.
  3. Tabla T/F por TEMPERATURA -> 'T. tend. cond. fase' (EDS nuevo).
  4. ESFUERZOS V/T/L de las hipotesis 1a-4a -> 'Cálculo de apoyos' y
     'Elección apoyos' (mismas ecuaciones que cambiar_cadena_conductor).
  5. Graficos de utilizacion (PNG/PDF) y fisica_linea.json (para el 3D).

GENERAL: funciona para cualquier linea/excel de la plantilla (lee el
conductor actual, las zonas, los grupos de vano de regulacion, etc.).

CLI:
  py editar_tension.py --listado p.json --resultado r.json
  py editar_tension.py --peticion p.json --resultado r.json
"""
import json
import os
import shutil

import openpyxl

from twinelec_paths import assets_dir

import mover_apoyo as mv
import tensiones_vanos as tv
import cambiar_cadena_conductor as ccc
from cambiar_cadena_conductor import (
    VIENTO_KMH,
    RUTA_APOYOS_CONFIG,
    RUTA_CONDUCTOR_CONFIG,
)


# ---------------------------------------------------------------------------
# Condiciones reglamentarias (clave, t_C, sobrecarga) -- igual que mover_apoyo
# ---------------------------------------------------------------------------
def _condiciones(ruta_excel):
    condiciones = mv.CONDICIONES_HIPOTESIS()
    if ccc._excel_tiene_viento_hielo(ruta_excel):
        condiciones.append(("TVHF", -15.0, "S_vh"))
    return condiciones


# ---------------------------------------------------------------------------
# Listado de tensiones actuales (para la ventana de Unity)
# ---------------------------------------------------------------------------
def listar_tensiones(ruta_excel):
    """Tensiones actuales de 'T. reg. cond. fase' por tramo, la lista de
    condiciones disponibles y el conductor vigente."""
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    try:
        ws = wb["T. reg. cond. fase"]
    except KeyError:
        raise ValueError("El Excel no tiene la hoja 'T. reg. cond. fase'.")
    col_tramo = ccc._localizar_columna(ws, [r"^tramo$"])
    condiciones = _condiciones(ruta_excel)
    etiquetas = {
        "TVF": r"^t maxima viento",
        "THF": r"^t maxima hielo",
        "TVHF": r"maxima.*hielo.*viento",
        "TVM": r"viento 1 2",
        "15V": r"^15",
        "T0H": r"^0",
        "TTemp": r"^50",
    }
    cols = {}
    for clave, patron in etiquetas.items():
        col = ccc._col_condicion_reg(ws, [patron])
        if col is not None:
            cols[clave] = col
    tramos = []
    for fila in ws.iter_rows(min_row=1):
        tramo = ccc._texto_limpio(fila[col_tramo].value) if col_tramo is not None else None
        if not tramo:
            continue
        nums = [int(n) for n in ccc.re.findall(r"\d+", tramo)]
        if len(nums) < 2:
            continue
        item = {"tramo": tramo, "i": nums[0], "j": nums[1], "tensiones": {}}
        for clave, col in cols.items():
            valor = ccc._numero_decimal(fila[col].value)
            if valor is not None:
                item["tensiones"][clave] = round(float(valor), 2)
        tramos.append(item)
    try:
        from datos_apoyos_excel import leer_conductor
        conductor = leer_conductor(ruta_excel).get("nomenclatura_excel")
    except Exception:
        conductor = None
    wb.close()
    return {
        "tramos": tramos,
        "condiciones": [c[0] for c in condiciones],
        "conductor": conductor,
        "t0_C": mv.leer_estado_eds(ruta_excel).get("t0_C"),
    }


# ---------------------------------------------------------------------------
# Edicion de una tension
# ---------------------------------------------------------------------------
def editar_tension(ruta_excel, salida_excel, tramo_i, tramo_j, condicion,
                   tension_daN):
    """Cambia la tension 'condicion' del tramo (tramo_i, tramo_j) y recalcula
    el resto de tensiones, la tabla por temperatura y los esfuerzos V/T/L."""
    if condicion not in {c[0] for c in _condiciones(ruta_excel)}:
        raise ValueError(
            "Condicion '%s' no disponible en esta linea (el Excel manda)."
            % condicion)
    tension_nueva = float(tension_daN)
    if tension_nueva <= 0:
        raise ValueError("La tension debe ser un numero positivo (daN).")

    if salida_excel != ruta_excel and not os.path.exists(salida_excel):
        shutil.copy2(ruta_excel, salida_excel)

    # --- Datos base -------------------------------------------------------
    conductor = ccc._leer_conductor_actual(ruta_excel)
    props = tv.propiedades_conductor(conductor)
    alfa = props["alfa_1_C"]
    E = props["E_daN_mm2"]
    S = props["seccion_mm2"]
    p0 = props["peso_daN_m"]

    geom = mv.leer_geometria_excel(ruta_excel)
    vanos = [t[0] for t in geom["tramos"]]
    desniveles = [t[1] for t in geom["tramos"]]
    n = geom["n"]
    if not (1 <= tramo_i <= n) or tramo_i + 1 != tramo_j:
        raise ValueError("Tramo no valido: debe ser (i, i+1) dentro de la linea.")
    idx = tramo_i - 1  # indice 0-based del tramo

    sobre = ccc.sobrecargas_por_tramo(ruta_excel, conductor, VIENTO_KMH)

    estado_eds = mv.leer_estado_eds(ruta_excel)
    t0 = estado_eds.get("t0_C", 10.0)
    estado_eds["p0_daN_m"] = p0
    t0_original = [estado_eds["T0_por_tramo"].get((i + 1, i + 2))
                   for i in range(n)]

    grupos_clave = mv.leer_grupos_vanos_regulacion(ruta_excel)
    grupos_por_indice = {}
    grupos_tramos = {}
    for i in range(n):
        g = grupos_clave.get((i + 1, i + 2), i + 1)
        grupos_por_indice[i] = g
        grupos_tramos.setdefault(g, []).append(i)

    grupo_editado = grupos_por_indice[idx]
    idxs_grupo = grupos_tramos[grupo_editado]
    ar = tv.vano_regulacion([vanos[i] for i in idxs_grupo])
    if ar is None or ar <= 0:
        raise ValueError("Vano de regulacion no valido para el tramo.")

    # --- Inversa del cambio de estado -------------------------------------
    condiciones = _condiciones(ruta_excel)
    pareja = next(c for c in condiciones if c[0] == condicion)
    t_cond, clave_p = pareja[1], pareja[2]
    p_cond = sobre[idx].get(clave_p)
    if p_cond is None or p_cond <= 0:
        raise ValueError("No hay sobrecarga '%s' para el tramo %d-%d."
                         % (clave_p, tramo_i, tramo_j))

    t0_nuevo = tv.t0_desde_tension_condicion(
        t0, tension_nueva, p_cond, t_cond, ar, p0, alfa, E, S)
    if t0_nuevo is None or t0_nuevo <= 0:
        raise ValueError(
            "No se pudo deducir un estado EDS valido para esa tension.")

    # --- Nuevo estado EDS (solo cambia el grupo editado) -------------------
    t0_nuevos = list(t0_original)
    for i in idxs_grupo:
        t0_nuevos[i] = t0_nuevo

    estado_nuevo = {
        "t0_C": t0,
        "T0_por_tramo": t0_nuevos,
        "p0_daN_m": p0,
    }
    tensiones = mv.recalcular_tensiones_tramos(
        vanos, desniveles, estado_nuevo, alfa, E, S, sobre,
        grupos=grupos_por_indice, condiciones=condiciones)

    # Valor real resultante en la condicion editada (comprobacion)
    par_editado = tensiones.get(idx, {}).get(condicion)
    real = par_editado["T"] if par_editado else None

    # --- Escritura en el Excel de trabajo ---------------------------------
    datos = {
        "tensiones": tensiones,
        "vanos": vanos,
        "desniveles": desniveles,
        "n": n,
        "t0_C": t0,
        "t0_por_tramo": t0_nuevos,
        "sobrecargas": sobre,
        "tiene_viento_hielo": ccc._excel_tiene_viento_hielo(ruta_excel),
    }
    wb = openpyxl.load_workbook(salida_excel)
    try:
        ccc.escribir_tensiones_reg(wb, datos)
        ccc.escribir_tabla_temperaturas(ruta_excel, wb, datos, conductor)
        wb.save(salida_excel)
    finally:
        wb.close()

    # --- Esfuerzos V/T/L, desviacion, flecha maxima, graficos ---------------
    ccc.recalcular_esfuerzos_apoyos(
        salida_excel, salida_excel, tensiones=tensiones, conductor=conductor)

    # --- fisica_linea.json (para que el 3D refleje la nueva tension) --------
    regenerada = _regenerar_fisica_linea(salida_excel, conductor)

    return {
        "ok": True,
        "tramo": "%d-%d" % (tramo_i, tramo_j),
        "condicion": condicion,
        "tension_pedida_daN": round(tension_nueva, 2),
        "tension_resultante_daN": round(real, 2) if real else None,
        "t0_anterior_daN": round(t0_original[idx], 2) if t0_original[idx] else None,
        "t0_nuevo_daN": round(t0_nuevo, 2),
        "grupo_vanos": grupo_editado,
        "tramos_grupo": [i + 1 for i in idxs_grupo],
        "vanos_regulacion_m": round(ar, 2),
        "fisica_linea_regenerada": regenerada,
        "salida_excel": os.path.abspath(salida_excel),
    }

# ---------------------------------------------------------------------------
# Regeneracion de fisica_linea.json (3D) -- best effort
# ---------------------------------------------------------------------------
def _regenerar_fisica_linea(ruta_excel, conductor):
    try:
        import principal
        from generar_doc_utm import leer_utm_actual
        carpeta_assets = str(assets_dir())
        ruta_carpeta = os.path.dirname(os.path.abspath(ruta_excel))
        datos_crudos = leer_utm_actual(ruta_carpeta)
        if len(datos_crudos) < 2:
            return False
        os.chdir(ruta_carpeta)
        principal.procesar_datos_fisicos(
            carpeta_assets, datos_crudos, conductor, archivo=ruta_excel)
        print("  [gemelo] fisica_linea.json regenerada con la nueva tension.")
        return True
    except Exception as error:
        print("  ⚠️ No se pudo regenerar fisica_linea.json:", error)
        return False


# ---------------------------------------------------------------------------
# CLI (JSON) para Unity
# ---------------------------------------------------------------------------
def procesar_peticion_json(ruta_peticion, ruta_resultado=None):
    with open(ruta_peticion, encoding="utf-8-sig") as archivo:
        p = json.load(archivo)
    try:
        excel = p["excel"]
        salida = p.get("salida_excel") or excel
        resultado = editar_tension(
            excel, salida,
            int(p["tramo_i"]), int(p["tramo_j"]),
            p["condicion"], float(p["tension_daN"]))
    except Exception as error:
        resultado = {"ok": False, "error": str(error)}
    if ruta_resultado:
        with open(ruta_resultado, "w", encoding="utf-8") as archivo:
            json.dump(resultado, archivo, ensure_ascii=False, indent=2,
                      default=str)
    return resultado


def procesar_listado_json(ruta_peticion, ruta_resultado=None):
    with open(ruta_peticion, encoding="utf-8-sig") as archivo:
        p = json.load(archivo)
    try:
        resultado = {"ok": True, "datos": listar_tensiones(p["excel"])}
    except Exception as error:
        resultado = {"ok": False, "error": str(error)}
    if ruta_resultado:
        with open(ruta_resultado, "w", encoding="utf-8") as archivo:
            json.dump(resultado, archivo, ensure_ascii=False, indent=2,
                      default=str)
    return resultado


if __name__ == "__main__":
    import sys
    import io
    try:
        if sys.stdout.encoding and "utf" not in sys.stdout.encoding.lower():
            sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                          errors="replace")
    except Exception:
        pass

    if "--listado" in sys.argv:
        idx = sys.argv.index("--listado")
        ruta_peticion = sys.argv[idx + 1]
        ruta_resultado = None
        if "--resultado" in sys.argv:
            ruta_resultado = sys.argv[sys.argv.index("--resultado") + 1]
        print(json.dumps(procesar_listado_json(ruta_peticion, ruta_resultado),
                         ensure_ascii=False, indent=2, default=str))
    elif "--peticion" in sys.argv:
        idx = sys.argv.index("--peticion")
        ruta_peticion = sys.argv[idx + 1]
        ruta_resultado = None
        if "--resultado" in sys.argv:
            ruta_resultado = sys.argv[sys.argv.index("--resultado") + 1]
        print(json.dumps(procesar_peticion_json(ruta_peticion, ruta_resultado),
                         ensure_ascii=False, indent=2, default=str))
    else:
        print("Uso: py editar_tension.py --listado p.json [--resultado r.json]")
        print("     py editar_tension.py --peticion p.json [--resultado r.json]")
