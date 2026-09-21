# -*- coding: utf-8 -*-
"""
cambiar_cruceta.py
==================
Motor de edicion de la SEMICRUCETA de un apoyo: cambiar su longitud
(1,25 / 1,50 / 1,75 / 2,00 / 2,25 / 2,50 m) y su tipo
(ASC = semicruceta plana, ATC = semicruceta atirantada), o sustituirla
por otra del catalogo. El catalogo se lee del propio script de Blender
(BlenderScripts/crear_semicrucetas_andel.py), igual que principal.py.

Cuando se cambia la semicruceta de un apoyo:
  1. Se actualizan las columnas del Excel 'Apoyos y crucetas':
       - 'Longitud crucetas m'   (a = semilongitud, en metros)
       - 'Referencia cruceta'
       - 'Cruceta tipo'
  2. Se actualiza apoyos_configurados.json (modelo_id, a_m, es_atirantada,
     referencia_origen, configuracion_id) para que Unity reconstruya el
     apoyo con el modelo nuevo.
  3. NO cambian ni HREF, ni desniveles, ni tensiones: la altura de amarre
     de los cables no depende de la longitud de la semicruceta.

Solo semicrucetas (ASC/ATC) por ahora; los armados se quedan para despues.

GENERAL: funciona para cualquier excel de la plantilla (la referencia se
comprueba contra el catalogo 3D real).

USO:
    from cambiar_cruceta import cambiar_cruceta_apoyo
    res = cambiar_cruceta_apoyo("line_example.xlsx", apoyo_numero=2,
                                nuevo_tipo="ASC", nueva_longitud_m=2.0,
                                salida_excel="line_example_editada.xlsx")

CLI (para Unity):
    py cambiar_cruceta.py --peticion peticion_cruceta.json --resultado resultado_cruceta.json
"""
import ast
import os
import re
import json

import openpyxl

from twinelec_paths import assets_dir, blender_scripts_dir

from datos_apoyos_excel import (
    _clave_texto,
    _localizar_columna,
    _numero_decimal,
    _texto_limpio,
)

# Tipos de semicruceta soportados (por ahora solo semicrucetas)
TIPOS_SEMICRUCETA = ("ASC", "ATC")

# Longitudes normalizadas (codigo -> metros) de la Serie C
CODIGO_LONGITUD = {
    "12": 1.25, "15": 1.50, "17": 1.75, "20": 2.00, "22": 2.25, "25": 2.50,
}

# Compatibilidad semicruceta <-> montaje (mismas reglas que Unity
# MontarCrucetasAndel: TRESBOLILLO/BANDERA admiten ASC y ATC; doble circuito
# solo atirantada).
COMPATIBILIDAD_MONTAJE = {
    "ASC": {"TRESBOLILLO", "BANDERA"},
    "ATC": {"TRESBOLILLO", "BANDERA", "DOBLE_CIRCUITO"},
}


# ---------------------------------------------------------------------------
# Catalogo 3D de semicrucetas (leido del script de Blender, como principal.py)
# ---------------------------------------------------------------------------
def _cargar_tablas_catalogo(ruta_script=None):
    """Lee las tablas de referencias del catalogo de Blender con 'ast'."""
    if ruta_script is None:
        ruta_script = str(blender_scripts_dir() / "crear_semicrucetas_andel.py")
    if not os.path.exists(ruta_script):
        raise FileNotFoundError("No se encontro el catalogo 3D: %s" % ruta_script)
    with open(ruta_script, "r", encoding="utf-8-sig") as f:
        arbol = ast.parse(f.read(), filename=ruta_script)
    nombres = {
        "ALTURAS_ACOPLE", "SEMICRUCETAS_ASC", "CRUCETAS_ARCX", "CRUCETAS_AMC",
        "SEMICRUCETAS_ATIRANTADAS", "SERIE_C_SEMICRUCETAS",
        "SERIE_C_CRUCETAS_RECTAS", "SERIE_C_ATIRANTADAS",
        "SERIE_C_BOVEDAS_TRIANGULO", "SERIE_C_BOVEDAS_CAPA",
        "SERIE_C_BOVEDAS_PICO", "MONTAJES_TRIANGULO_TG",
        "BOVEDAS_CAPA", "BOVEDAS_PICO", "BOVEDAS_TRIANGULO",
        "CATALOGO_13000_CRUCETAS", "PRESILLA_TRESBOLILLO",
        "PRESILLA_RECTAS_MONTAJE_0", "PRESILLA_RECTAS_ARRIOSTRADAS",
        "PRESILLA_BOVEDAS_PICO", "GANCHO_HERRAJE_PUENTE",
    }
    tablas = {}
    for nodo in arbol.body:
        if not isinstance(nodo, ast.Assign) or len(nodo.targets) != 1:
            continue
        destino = nodo.targets[0]
        if isinstance(destino, ast.Name) and destino.id in nombres:
            tablas[destino.id] = ast.literal_eval(nodo.value)
    return tablas


def _modelos_catalogados(tablas):
    """Conjunto de modelo_id validos (misma logica que principal.py:
    cargar_modelos_catalogados, incluidas las variantes H de acople)."""
    modelos = set(tablas.get("SEMICRUCETAS_ASC", {}))
    modelos.update(tablas.get("CRUCETAS_ARCX", {}))
    modelos.update(tablas.get("CRUCETAS_AMC", {}))
    modelos.add(tablas.get("GANCHO_HERRAJE_PUENTE"))
    modelos.update(tablas.get("MONTAJES_TRIANGULO_TG", {}))
    alturas = [sufijo for sufijo, _, _ in tablas.get("ALTURAS_ACOPLE", [])]
    for nombre_tabla in ("BOVEDAS_CAPA", "BOVEDAS_PICO", "BOVEDAS_TRIANGULO",
                         "SEMICRUCETAS_ATIRANTADAS"):
        for referencia in tablas.get(nombre_tabla, {}):
            modelos.update("%s_%s" % (referencia, altura) for altura in alturas)
    for nombre_tabla in ("SERIE_C_SEMICRUCETAS", "SERIE_C_CRUCETAS_RECTAS",
                         "SERIE_C_ATIRANTADAS", "SERIE_C_BOVEDAS_TRIANGULO",
                         "SERIE_C_BOVEDAS_CAPA", "SERIE_C_BOVEDAS_PICO"):
        for referencia in tablas.get(nombre_tabla, {}):
            modelos.add("C_%s" % referencia)
    for referencia in tablas.get("CATALOGO_13000_CRUCETAS", {}):
        modelos.add("CAT13000_%s" % referencia)
    modelos.add(tablas.get("PRESILLA_TRESBOLILLO"))
    modelos.update(tablas.get("PRESILLA_RECTAS_MONTAJE_0", {}))
    modelos.update(tablas.get("PRESILLA_RECTAS_ARRIOSTRADAS", {}))
    for referencia, _, _ in tablas.get("PRESILLA_BOVEDAS_PICO", []):
        modelos.add(referencia)
    modelos.discard(None)
    return modelos


def opciones_semicrucetas():
    """[(tipo, codigo, longitud_m), ...] de semicrucetas ASC y ATC del
    catalogo, ordenadas por tipo y longitud. Es la lista que mostrara Unity."""
    tablas = _cargar_tablas_catalogo()
    opciones = []
    for tipo, tabla in (("ASC", "SERIE_C_SEMICRUCETAS"),
                        ("ATC", "SERIE_C_ATIRANTADAS")):
        for referencia, longitud in tablas.get(tabla, {}).items():
            m = re.search(r"(\d+)\s*$", referencia)
            if m:
                opciones.append((tipo, m.group(1), float(longitud)))
    opciones.sort(key=lambda o: (o[0], o[2]))
    return opciones


# ---------------------------------------------------------------------------
# Lectura de la semicruceta actual y del montaje desde el Excel
# ---------------------------------------------------------------------------
def _fila_de_apoyo(ws, apoyo_numero):
    for fila in ws.iter_rows(min_row=1, max_col=2):
        numero = _numero_decimal(fila[0].value)
        if numero is not None and int(numero) == int(apoyo_numero):
            return fila[0].row
    return None


def leer_cruceta_excel(ruta_excel):
    """Devuelve {apoyo_numero: {'a_m','referencia','tipo'}} desde la pestana
    'Apoyos y crucetas'."""
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws = wb["Apoyos y crucetas"]
    col_a = _localizar_columna(ws, [r"longitud.*crucetas"], fila_cabecera=2)
    col_ref = _localizar_columna(ws, [r"referencia.*cruceta"], fila_cabecera=2)
    col_tipo = _localizar_columna(ws, [r"cruceta.*tipo"], fila_cabecera=2)
    resultado = {}
    for fila in ws.iter_rows(min_row=1):
        numero = _numero_decimal(fila[0].value)
        if numero is None:
            continue
        resultado[int(numero)] = {
            "a_m": (_numero_decimal(fila[col_a].value)
                    if col_a is not None else None),
            "referencia": (_texto_limpio(fila[col_ref].value)
                           if col_ref is not None else None),
            "tipo": (_texto_limpio(fila[col_tipo].value)
                     if col_tipo is not None else None),
        }
    return resultado


def _normalizar_montaje(texto):
    clave = _clave_texto(texto)
    if not clave:
        return None
    if "tresbolillo" in clave:
        return "TRESBOLILLO"
    if "doble circuito" in clave:
        return "DOBLE_CIRCUITO"
    if "bandera" in clave:
        return "BANDERA"
    if "horizontal" in clave:
        return "HORIZONTAL"
    if "paso" in clave and "lateral" in clave:
        return "PASO_LATERAL"
    if "paso" in clave and "herraje" in clave:
        return "PASO_HERRAJE_PUENTE"
    if "triangulo" in clave or re.search(r"\btg\s*\d+", clave):
        return "TRIANGULO_TG"
    if "boveda" in clave and "pico" in clave:
        return "BOVEDA_PICO"
    if "boveda" in clave and "capa" in clave:
        return "BOVEDA_CAPA"
    if "boveda" in clave and "triangulo" in clave:
        return "BOVEDA_TRIANGULO"
    return None


def leer_montaje_por_apoyo(ruta_excel):
    """Devuelve {apoyo_numero: tipo_montaje_normalizado} desde 'Armado base'."""
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws = wb["Apoyos y crucetas"]
    col = _localizar_columna(ws, [r"armado.*base"], fila_cabecera=2)
    resultado = {}
    for fila in ws.iter_rows(min_row=1):
        numero = _numero_decimal(fila[0].value)
        if numero is None:
            continue
        resultado[int(numero)] = _normalizar_montaje(
            _texto_limpio(fila[col].value) if col is not None else None)
    return resultado


# ---------------------------------------------------------------------------
# Escritura en apoyos_configurados.json
# ---------------------------------------------------------------------------
def _reconstruir_config_id(apoyo_obj):
    """Reconstruye el configuracion_id con el mismo formato de principal.py
    (_crear_configuracion_id): CFG1__catalogo__armado__montaje__cruceta__
    A####__SEP####__H####__REC####__CUP####."""
    partes = [
        "CFG1",
        apoyo_obj.get("catalogo_id") or "CATALOGO_DESCONOCIDO",
        (apoyo_obj.get("armado") or {}).get("modelo_id") or "ARMADO_NO_RESUELTO",
        (apoyo_obj.get("montaje") or {}).get("tipo") or "MONTAJE_NO_RESUELTO",
        (apoyo_obj.get("cruceta") or {}).get("modelo_id") or "MODELO_NO_RESUELTO",
    ]
    dimensiones = (
        ("A", (apoyo_obj.get("cruceta") or {}).get("a_m")),
        ("SEP", (apoyo_obj.get("montaje") or {}).get("separacion_vertical_m")),
        ("H", (apoyo_obj.get("apoyo") or {}).get("altura_total_m")),
        ("REC", (apoyo_obj.get("apoyo") or {}).get("recrecido_cabeza_m")),
        ("CUP", (apoyo_obj.get("montaje") or {}).get("altura_cupula_m")),
    )
    for etiqueta, valor in dimensiones:
        if valor is not None:
            partes.append("%s%d" % (etiqueta, int(round(float(valor) * 1000))))
    identificador = "__".join(partes).upper()
    return re.sub(r"[^A-Z0-9_-]+", "_", identificador)


def actualizar_config_unity(ruta_config, apoyo_numero, modelo_id, tipo_origen,
                            referencia_origen, a_m, es_atirantada):
    """Actualiza la cruceta de un apoyo en apoyos_configurados.json y
    reconstruye su configuracion_id."""
    with open(ruta_config, encoding="utf-8-sig") as f:
        cfg = json.load(f)
    cambiado = False
    for apoyo_obj in cfg.get("apoyos", []):
        apoyo = apoyo_obj.get("apoyo") or {}
        if apoyo.get("numero") != int(apoyo_numero):
            continue
        cruceta = apoyo_obj.setdefault("cruceta", {})
        cruceta["modelo_id"] = modelo_id
        cruceta["tipo_origen"] = tipo_origen
        cruceta["referencia_origen"] = referencia_origen
        cruceta["a_m"] = round(float(a_m), 3)
        cruceta["es_atirantada"] = bool(es_atirantada)
        apoyo_obj["configuracion_id"] = _reconstruir_config_id(apoyo_obj)
        cambiado = True
        break
    if not cambiado:
        raise ValueError("No se encontro el apoyo %s en '%s'."
                         % (apoyo_numero, ruta_config))
    with open(ruta_config, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    return ruta_config


# ---------------------------------------------------------------------------
# Orquestador principal
# ---------------------------------------------------------------------------
def _ruta_config_unity(carpeta_excel):
    """Ruta al estado compartido de apoyos que consume Unity."""
    return str(assets_dir() / "apoyos_configurados.json")


def cambiar_cruceta_apoyo(ruta_excel, apoyo_numero, nuevo_tipo=None,
                          nueva_longitud_m=None, nueva_referencia=None,
                          salida_excel=None, config_unity=None):
    """Cambia la semicruceta de un apoyo por otra del catalogo.

    Se puede indicar nuevo_tipo ("ASC"/"ATC") + nueva_longitud_m, o
    nueva_referencia ("ASC-20", "C_ATC-15"...). Actualiza el Excel de
    trabajo y apoyos_configurados.json. Devuelve un dict resumen."""
    # ---- Normalizar la peticion a (tipo, codigo, longitud) ----
    if nueva_referencia:
        ref = str(nueva_referencia).strip().upper().replace(" ", "")
        ref = re.sub(r"^C_", "", ref)
        m = re.search(r"^(ASC|ATC)\s*[-]?\s*(\d+)$", ref)
        if not m:
            raise ValueError(
                "Referencia de semicruceta no valida: %r. Usa por ejemplo "
                "'ASC-20' o 'ATC-15'." % (nueva_referencia,))
        nuevo_tipo, codigo = m.group(1), m.group(2)
        nueva_longitud_m = CODIGO_LONGITUD.get(codigo)
        if nueva_longitud_m is None:
            raise ValueError("Codigo de longitud desconocido: %s." % codigo)

    nuevo_tipo = str(nuevo_tipo or "").upper().strip()
    if nuevo_tipo not in TIPOS_SEMICRUCETA:
        raise ValueError(
            "Tipo de semicruceta no soportado: %r. De momento solo "
            "ASC (plana) o ATC (atirantada)." % nuevo_tipo)

    if nueva_longitud_m is None:
        raise ValueError("Falta 'nueva_longitud_m' (1.25, 1.5, 1.75, 2.0, "
                         "2.25 o 2.5 m).")
    try:
        nueva_longitud_m = float(nueva_longitud_m)
    except (TypeError, ValueError):
        raise ValueError("nueva_longitud_m debe ser un numero: %r"
                         % (nueva_longitud_m,))

    codigo = next((c for c, l in CODIGO_LONGITUD.items()
                   if abs(l - nueva_longitud_m) < 1e-9), None)
    if codigo is None:
        raise ValueError(
            "Longitud %.2f m no esta en el catalogo (12/15/17/20/22/25)."
            % nueva_longitud_m)

    # ---- Comprobar el modelo contra el catalogo 3D real ----
    tablas = _cargar_tablas_catalogo()
    modelos = _modelos_catalogados(tablas)
    modelo_id = "C_%s-%s" % (nuevo_tipo, codigo)
    if modelo_id not in modelos:
        raise ValueError("El modelo %s no existe en el catalogo 3D."
                         % modelo_id)
    es_atirantada = nuevo_tipo == "ATC"
    nueva_referencia = "%s-%s" % (nuevo_tipo, codigo)

    # ---- Excel de trabajo (acumulativo) ----
    if salida_excel is None:
        base = os.path.splitext(os.path.basename(ruta_excel))[0]
        salida_excel = os.path.join(os.path.dirname(os.path.abspath(ruta_excel)),
                                    base + "_editada.xlsx")

    # RECIPROCIDAD: leer el estado actual del editada si existe
    salida_abs = os.path.abspath(salida_excel)
    fuente = (salida_abs if (os.path.exists(salida_abs)
                             and salida_abs != os.path.abspath(ruta_excel))
              else ruta_excel)

    crucetas = leer_cruceta_excel(fuente)
    montajes = leer_montaje_por_apoyo(fuente)
    apoyo = int(apoyo_numero)
    actual = crucetas.get(apoyo)
    if actual is None:
        raise ValueError("No se encontraron datos de cruceta del apoyo %s."
                         % apoyo_numero)

    # ---- Compatibilidad con el montaje ----
    montaje = montajes.get(apoyo)
    permitidos = COMPATIBILIDAD_MONTAJE.get(nuevo_tipo, set())
    if montaje is not None and montaje not in permitidos:
        raise ValueError(
            "El apoyo %d tiene montaje %s, incompatible con semicrucetas %s. "
            "Compatibles: %s."
            % (apoyo, montaje, nuevo_tipo,
               ", ".join(sorted(permitidos)) if permitidos else "ninguno"))

    # ---- Escribir en el Excel ----
    import shutil
    if not os.path.exists(salida_excel):
        os.makedirs(os.path.dirname(salida_abs), exist_ok=True)
        shutil.copy2(ruta_excel, salida_excel)
        print("  [cambiar_cruceta] Copia creada:", salida_excel)
    wb = openpyxl.load_workbook(salida_excel)
    ws = wb["Apoyos y crucetas"]
    col_a = _localizar_columna(ws, [r"longitud.*crucetas"], fila_cabecera=2)
    col_ref = _localizar_columna(ws, [r"referencia.*cruceta"], fila_cabecera=2)
    col_tipo = _localizar_columna(ws, [r"cruceta.*tipo"], fila_cabecera=2)
    fila = _fila_de_apoyo(ws, apoyo)
    if fila is None:
        raise ValueError("No se encontro la fila del apoyo %s en el Excel."
                         % apoyo_numero)
    if col_a is not None:
        ws.cell(row=fila, column=col_a + 1).value = round(nueva_longitud_m, 2)
    if col_ref is not None:
        ws.cell(row=fila, column=col_ref + 1).value = nueva_referencia
    if col_tipo is not None:
        ws.cell(row=fila, column=col_tipo + 1).value = nueva_referencia

    # Reciprocidad total: normaliza tambien las celdas del montaje
    # ('Armado base' y 'Referencia armado') al valor canonico del montaje actual.
    from cambiar_tipo_apoyo import MONTAJES as _MONTAJES_TIPO
    if montaje in _MONTAJES_TIPO:
        col_armado = _localizar_columna(ws, [r"armado.*base"], fila_cabecera=2)
        col_ref_armado = _localizar_columna(ws, [r"referenc.*armado"],
                                            fila_cabecera=2)
        if col_armado is not None:
            ws.cell(row=fila, column=col_armado + 1).value = \
                _MONTAJES_TIPO[montaje]["texto_excel"]
        if col_ref_armado is not None:
            ws.cell(row=fila, column=col_ref_armado + 1).value = \
                _MONTAJES_TIPO[montaje]["referencia_armado"]
    wb.save(salida_excel)

    # ---- Configuracion de Unity ----
    ruta_config = None
    if config_unity is None:
        config_unity = _ruta_config_unity(
            os.path.dirname(os.path.abspath(ruta_excel)))
    if os.path.exists(config_unity):
        ruta_config = actualizar_config_unity(
            config_unity, apoyo, modelo_id, nueva_referencia,
            nueva_referencia, nueva_longitud_m, es_atirantada)
        print("  [cambiar_cruceta] Unity actualizado:", ruta_config)

    return {
        "apoyo_numero": apoyo,
        "modelo_anterior": actual.get("tipo"),
        "modelo_nuevo": modelo_id,
        "referencia": nueva_referencia,
        "tipo": nuevo_tipo,
        "longitud_m": round(nueva_longitud_m, 3),
        "es_atirantada": es_atirantada,
        "montaje": montaje,
        "salida_excel": os.path.abspath(salida_excel),
        "config_unity": os.path.abspath(ruta_config) if ruta_config else None,
    }


# ---------------------------------------------------------------------------
# CLI (JSON) para Unity
# ---------------------------------------------------------------------------
def procesar_peticion_json(ruta_peticion, ruta_resultado=None):
    """Lee la peticion JSON de Unity, cambia la semicruceta y escribe el
    resultado (o el error). Peticion:
        {"excel":.., "utm":.., "salida_excel":.., "apoyo":N,
         "tipo":"ASC"|"ATC", "longitud":2.0, "referencia":"ASC-20", ...}
    Se puede dar 'tipo'+'longitud' o 'referencia'."""
    with open(ruta_peticion, encoding="utf-8-sig") as f:
        p = json.load(f)
    try:
        res = cambiar_cruceta_apoyo(
            p.get("excel"), int(p["apoyo"]),
            nuevo_tipo=p.get("tipo"),
            nueva_longitud_m=p.get("longitud"),
            nueva_referencia=p.get("referencia"),
            salida_excel=p.get("salida_excel"),
            config_unity=p.get("config_unity"))
        resultado = {"ok": True}
        resultado.update(res)
    except Exception as e:
        resultado = {"ok": False, "error": str(e)}
    if ruta_resultado:
        with open(ruta_resultado, "w", encoding="utf-8") as f:
            json.dump(resultado, f, ensure_ascii=False, indent=2, default=str)
    return resultado


if __name__ == "__main__":
    import sys
    import io
    try:
        if sys.stdout.encoding and "utf" not in sys.stdout.encoding.lower():
            sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                          errors="replace")
    except Exception:
        pass

    if "--peticion" in sys.argv:
        idx = sys.argv.index("--peticion")
        ruta_peticion = sys.argv[idx + 1]
        ruta_resultado = None
        if "--resultado" in sys.argv:
            ruta_resultado = sys.argv[sys.argv.index("--resultado") + 1]
        r = procesar_peticion_json(ruta_peticion, ruta_resultado)
        print(json.dumps(r, ensure_ascii=False, indent=2)[:2500])
    else:
        CARPETA = os.path.dirname(os.path.abspath(__file__))
        RUTA = os.path.join(CARPETA, "line_example.xlsx")
        SALIDA = os.path.join(CARPETA, "line_example_editada.xlsx")

        print("== DEMO: cambiar la semicruceta del apoyo 2 a ASC-20 (plana) ==")
        res = cambiar_cruceta_apoyo(RUTA, apoyo_numero=2,
                                    nuevo_tipo="ASC", nueva_longitud_m=2.0,
                                    salida_excel=SALIDA)
        print("  Anterior:", res["modelo_anterior"], "-> Nuevo:", res["modelo_nuevo"])
        print("  Tipo: %s | Longitud: %.2f m | Atirantada: %s"
              % (res["tipo"], res["longitud_m"], res["es_atirantada"]))
        print("  Excel de trabajo:", res["salida_excel"])
        print("  Unity actualizado:", res["config_unity"])

        print()
        print("== CATALOGO DE SEMICRUCETAS DISPONIBLE ==")
        for tipo, codigo, longitud in opciones_semicrucetas():
            print("   %s-%s  (%.2f m, %s)" % (tipo, codigo, longitud,
                                              "atirantada" if tipo == "ATC"
                                              else "plana"))


