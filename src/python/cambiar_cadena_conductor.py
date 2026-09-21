# -*- coding: utf-8 -*-
"""
cambiar_cadena_conductor.py
===========================
Cambia el tipo de CADENAS DE AISLADORES y/o el CONDUCTOR de la linea del
gemelo digital, y recalcula en cascada los valores que dependen de ellos:

  CADENAS  -> ambito "apoyo" (solo un apoyo) o "todos" (toda la linea).
              Catalogo: catalogo_cadenas_completo.json (Andelec v25).
  CONDUCTOR-> comun a toda la linea.
              Catálogo: data/catalogs/conductors/catalogo_conductores.json.

Al aplicar se actualiza:
  1. EXCEL:
     - 'Tipo de conductor' (hoja Mediciones) y columna 'Conductor' de
       'T. tend. cond. fase' y 'T. reg. cond. fase'.
     - Numero de cadenas por tipo (hoja Mediciones, filas de cadenas).
     - TENSIONES y FLECHAS del conductor (E.D.S. constante, motor de
       tensiones_vanos): 'T. reg. cond. fase' (hipotesis) y
       'T. tend. cond. fase' (tabla por temperatura).
     - DESVIACION DE CADENA y FLECHA MAXIMA de cada apoyo (col. AA y AB de
       'Cálculo de apoyos').
     - ESFUERZOS V/T/L de las hipotesis 1a-4a de cada apoyo (columnas F-W de
       'Cálculo de apoyos'), con las ecuaciones del Tema 10.
  2. UNITY:
     - apoyos_configurados.json  -> bloque 'cadena' por apoyo.
     - conductor_configurado.json-> propiedades del nuevo conductor (+ R DLR).

CLI:
  py cambiar_cadena_conductor.py --catalogo_cadenas [--conductor LA-56]
  py cambiar_cadena_conductor.py --catalogo_conductores
  py cambiar_cadena_conductor.py --peticion p.json --resultado r.json
  py cambiar_cadena_conductor.py --sincronizar p.json --resultado r.json
"""
import json
import math
import os
import re
import shutil
import unicodedata

import openpyxl

from twinelec_paths import assets_dir, catalogs_dir

import tensiones_vanos as tv
import mover_apoyo as mv
from sobrecargas import calcular_sobrecargas
from datos_apoyos_excel import (
    _clave_conductor,
    _localizar_columna,
    _numero_decimal,
    _texto_limpio,
    n_cadenas,
    mapear_tipo_a_clave,
    extraer_zonas_por_tramo,
)

# ---------------------------------------------------------------------------
# Rutas y constantes
# ---------------------------------------------------------------------------
BASE = os.path.dirname(os.path.abspath(__file__))
RUTA_CATALOGO_CADENAS = str(catalogs_dir() / "catalogo_cadenas_completo.json")
RUTA_CATALOGO_CONDUCTORES = str(
    catalogs_dir() / "conductors" / "catalogo_conductores.json"
)
RUTA_RESISTENCIAS = str(
    catalogs_dir() / "conductors" / "resistencias_electricas.json"
)
RUTA_APOYOS_CONFIG = str(assets_dir() / "apoyos_configurados.json")
RUTA_CONDUCTOR_CONFIG = str(assets_dir() / "conductor_configurado.json")

NUMERO_FASES = 3
VIENTO_KMH = 120.0  # linea no especial (<= 220 kV) -> viento reglamentario 120


# ---------------------------------------------------------------------------
# Utilidades de texto (misma convencion que mover_apoyo / datos_apoyos_excel)
# ---------------------------------------------------------------------------
def _clave_normalizada(valor):
    t = _texto_limpio(valor)
    if not t:
        return ""
    t = unicodedata.normalize("NFKD", t)
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", t.lower()).strip()


def _numero(valor):
    t = _texto_limpio(valor)
    if t is None:
        return None
    t = t.replace(",", ".").replace(" ", "")
    try:
        return float(t)
    except (TypeError, ValueError):
        return None


# ---------------------------------------------------------------------------
# Catalogos (para los desplegables de Unity)
# ---------------------------------------------------------------------------
def clasificar_cadena(nombre):
    """Extrae tipo (SUS/ANC), material (VID/POL) y simple/doble del nombre."""
    nombre_aux = (nombre or "").upper()
    tipo = "SUS" if "SUS" in nombre_aux else ("ANC" if "ANC" in nombre_aux else "")
    material = "VID" if "VID" in nombre_aux else ("POL" if "POL" in nombre_aux else "")
    doble = "DOB" in nombre_aux
    return {"tipo": tipo, "material": material, "doble": doble}


def cargar_catalogo_cadenas():
    with open(RUTA_CATALOGO_CADENAS, encoding="utf-8") as archivo:
        return json.load(archivo).get("cadenas", [])


def catalogo_cadenas(conductor_filtro=None):
    """Lista JSON-serializable de cadenas para el desplegable de Unity.

    Si 'conductor_filtro' se indica, se anade un campo 'compatible' para poder
    agrupar/ordenar las cadenas por el conductor de la linea.
    """
    clave_filtro = _clave_conductor(conductor_filtro) if conductor_filtro else None
    resultado = []
    for c in cargar_catalogo_cadenas():
        car = c.get("caracteristicas", {})
        resultado.append({
            "nombre": c["nombre"],
            "denominacion": c.get("denominacion", ""),
            "conductor": c.get("conductor"),
            "tension_kv": c.get("tension_kv"),
            "tipo": clasificar_cadena(c["nombre"])["tipo"],
            "material": clasificar_cadena(c["nombre"])["material"],
            "doble": clasificar_cadena(c["nombre"])["doble"],
            "longitud_m": car.get("longitud_m"),
            "peso_daN": car.get("peso_daN"),
            "carga_rotura_daN": car.get("carga_rotura_daN"),
            "compatible": (clave_filtro is None
                           or (_clave_conductor(c.get("conductor")) == clave_filtro)),
        })
    return resultado


def buscar_cadena(nombre):
    for c in cargar_catalogo_cadenas():
        if c["nombre"] == nombre:
            return c
    return None


def cargar_catalogo_conductores():
    with open(RUTA_CATALOGO_CONDUCTORES, encoding="utf-8") as archivo:
        return json.load(archivo)


def catalogo_conductores():
    """Lista JSON-serializable de conductores para el desplegable de Unity."""
    catalogo = cargar_catalogo_conductores()
    resultado = []
    for c in catalogo.get("conductores", []):
        aliases = c.get("aliases", [])
        nomenclatura = next((a for a in aliases if re.match(r"^[A-Z]+\s*-\s*\d", a)),
                            c["designacion"])
        resultado.append({
            "nomenclatura": nomenclatura,
            "designacion": c["designacion"],
            "composicion": c.get("composicion", ""),
            "diametro_mm": c["diametro_mm"],
            "seccion_mm2": c["seccion_mm2"],
            "masa_kg_km": c["masa_kg_km"],
        })
    return resultado


def resolver_conductor(nomenclatura):
    """Resuelve la nomenclatura (p.ej. 'LA-56') contra el catalogo de
    conductores y devuelve las propiedades mecanicas completas."""
    catalogo = cargar_catalogo_conductores()
    clave_buscada = _clave_conductor(nomenclatura)
    coincidencia = None
    for conductor in catalogo.get("conductores", []):
        nombres = [conductor["designacion"], *conductor.get("aliases", [])]
        if clave_buscada in {_clave_conductor(n) for n in nombres}:
            coincidencia = conductor
            break
    if coincidencia is None:
        raise ValueError(
            "El conductor '%s' no figura en el catalogo de conductores."
            % nomenclatura)
    props_composicion = {
        e["composicion"]: e
        for e in catalogo.get("propiedades_mecanicas_por_composicion", [])
    }
    composicion = coincidencia.get("composicion")
    propiedades = props_composicion.get(composicion, coincidencia)
    if "modulo_elasticidad_daN_mm2" not in propiedades:
        raise ValueError(
            "No hay modulo de elasticidad para la composicion '%s'." % composicion)
    if "coef_dilatacion_1_C" not in propiedades:
        raise ValueError(
            "No hay coeficiente de dilatacion para la composicion '%s'." % composicion)
    return {
        "nomenclatura_excel": nomenclatura,
        "designacion": coincidencia["designacion"],
        "composicion": composicion,
        "diametro_mm": coincidencia["diametro_mm"],
        "diametro_m": coincidencia["diametro_mm"] / 1000.0,
        "seccion_mm2": coincidencia["seccion_mm2"],
        "masa_kg_km": coincidencia["masa_kg_km"],
        "peso_kg_m": coincidencia["masa_kg_km"] / 1000.0,
        "peso_daN_m": coincidencia["masa_kg_km"] * 0.98 / 1000.0,
        "modulo_elasticidad_daN_mm2": propiedades["modulo_elasticidad_daN_mm2"],
        "coef_dilatacion_1_C": propiedades["coef_dilatacion_1_C"],
        "fuente_propiedades": coincidencia.get(
            "fuente_propiedades",
            "Tema 1 - Conductores - 2015.pdf, tabla 5, pagina 11"),
    }


# ---------------------------------------------------------------------------
# Recalculo de tensiones y flechas con el nuevo conductor (EDS constante)
# ---------------------------------------------------------------------------
def _fila_tramo_en(ws, col_tramo, tramo_i, tramo_j):
    """Fila (1-based) donde la columna 'Tramo' contiene '  i-  j'."""
    for fila in ws.iter_rows(min_row=1):
        v = _texto_limpio(fila[col_tramo].value) if col_tramo is not None else None
        if not v:
            continue
        nums = [int(n) for n in re.findall(r"\d+", v)]
        if len(nums) >= 2 and nums[0] == tramo_i and nums[1] == tramo_j:
            return fila[col_tramo].row
    return None


def sobrecargas_por_tramo(ruta_excel, conductor, viento_kmh=VIENTO_KMH):
    """Sobrecargas [daN/m] por tramo (Sv, Sv120, Sh, S_vm, Peso cable)."""
    geom = mv.leer_geometria_excel(ruta_excel)
    n = geom["n"]
    zonas = extraer_zonas_por_tramo(ruta_excel)
    p0 = conductor["peso_daN_m"]
    peso_kg_m = conductor["masa_kg_km"] / 1000.0
    diam_mm = conductor["diametro_mm"]
    lista = []
    for i in range(n):
        zona = zonas.get((i + 1, i + 2), "A")
        sob = calcular_sobrecargas(diam_mm, peso_kg_m, zona, viento_kmh=viento_kmh)
        lista.append({
            "Sv": sob["Sv"], "Sv120": sob["Sv120"], "Sh": sob["Sh"],
            "S_vm": sob["S_vm"], "Peso cable": p0,
            "zona": zona, "P_v_max": sob["P_v_max"], "F_v_120": sob["F_v_120"],
            "V_vh": sob["V_vh"], "S_vh": sob["S_vh"],
            "dmh": sob["dmh_m"],
        })
    return lista


def _excel_tiene_viento_hielo(ruta_excel):
    """True si 'T. reg. cond. fase' tiene la columna de viento+hielo CON
    valores. El programa opera sobre lo que trae el Excel: si la linea no
    incluye viento-hielo, no se calcula ni se escribe esa hipotesis."""
    try:
        wb = openpyxl.load_workbook(ruta_excel, data_only=True)
        try:
            ws = wb["T. reg. cond. fase"]
        except KeyError:
            return False
        col = None
        for c in ws[2]:
            clave = _clave_normalizada(c.value)
            if clave and "viento" in clave and "hielo" in clave:
                col = c.column - 1
                break
        if col is None:
            return False
        for fila in ws.iter_rows(min_row=4, max_col=col + 1):
            if _numero_decimal(fila[col].value) is not None:
                return True
        return False
    except Exception:
        return True


def recalcular_tensiones_linea(ruta_excel, conductor, viento_kmh=VIENTO_KMH):
    """Recalcula tensiones/flechas de TODA la linea con el nuevo conductor
    (E.D.S. constante, opcion A). Devuelve un dict con 'tensiones'
    {i_tramo: {clave: {'T','F'}}}, 'vanos', 'desniveles', 't0_C',
    't0_por_tramo' y 'sobrecargas'."""
    geom = mv.leer_geometria_excel(ruta_excel)
    vanos = [t[0] for t in geom["tramos"]]
    desniveles = [t[1] for t in geom["tramos"]]
    n = geom["n"]

    props = tv.propiedades_conductor(conductor)
    alfa = props["alfa_1_C"]
    E = props["E_daN_mm2"]
    S = props["seccion_mm2"]
    p0 = props["peso_daN_m"]

    sobre = sobrecargas_por_tramo(ruta_excel, conductor, viento_kmh)

    estado_eds = mv.leer_estado_eds(ruta_excel)
    estado_eds["p0_daN_m"] = p0
    t0s = []
    for i in range(n):
        t0s.append(estado_eds["T0_por_tramo"].get((i + 1, i + 2)))
    estado_eds["T0_por_tramo"] = t0s

    grupos_clave = mv.leer_grupos_vanos_regulacion(ruta_excel)
    grupos_por_indice = {}
    for i in range(n):
        grupos_por_indice[i] = grupos_clave.get((i + 1, i + 2), i + 1)

    condiciones = mv.CONDICIONES_HIPOTESIS()
    tiene_vh = _excel_tiene_viento_hielo(ruta_excel)
    if tiene_vh:
        condiciones.append(("TVHF", -15.0, "S_vh"))
    tensiones = mv.recalcular_tensiones_tramos(
        vanos, desniveles, estado_eds, alfa, E, S, sobre,
        grupos=grupos_por_indice, condiciones=condiciones)

    return {
        "tensiones": tensiones,
        "vanos": vanos,
        "desniveles": desniveles,
        "n": n,
        "t0_C": estado_eds.get("t0_C", 10.0),
        "t0_por_tramo": t0s,
        "sobrecargas": sobre,
        "tiene_viento_hielo": tiene_vh,
    }


def escribir_conductor_en_excel(wb, nomenclatura):
    """Actualiza 'Tipo de conductor' (Mediciones) y la columna 'Conductor' de
    'T. tend. cond. fase' y 'T. reg. cond. fase'."""
    ws_med = wb["Mediciones"]
    for fila in ws_med.iter_rows():
        if any("tipo de conductor" in _clave_normalizada(c.value) for c in fila):
            destino = None
            for c in fila:
                if _texto_limpio(c.value):
                    destino = c
            if destino is not None and destino.column > 1:
                destino.value = nomenclatura
            break

    for hoja in ("T. tend. cond. fase", "T. reg. cond. fase"):
        ws = wb[hoja]
        col_tramo = _localizar_columna(ws, [r"^tramo$"])
        col_cond = _localizar_columna(ws, [r"^conductor$"])
        if col_tramo is None or col_cond is None:
            continue
        for fila in ws.iter_rows(min_row=1):
            if _texto_limpio(fila[col_tramo].value) is not None:
                fila[col_cond].value = nomenclatura


def _col_condicion_reg(ws, patrones):
    """Columna (0-based) de 'T. reg. cond. fase' cuya cabecera de fila 2
    coincide (misma heuristica que mover_apoyo._columna_condicion_reg)."""
    for celda in ws[2]:
        clave = _clave_normalizada(celda.value)
        if clave and any(re.search(p, clave) for p in patrones):
            return celda.column - 1
    return None


def _tiene_flecha(ws, col_t):
    celda = ws.cell(row=3, column=col_t + 2).value
    return celda is not None and "f" in str(celda).lower()


def escribir_tensiones_reg(wb, datos):
    """Escribe T/F de las hipotesis en 'T. reg. cond. fase'."""
    ws = wb["T. reg. cond. fase"]
    col_tramo = _localizar_columna(ws, [r"^tramo$"])
    cols_cond = {
        "TVF": _col_condicion_reg(ws, [r"^t maxima viento"]),
        "THF": _col_condicion_reg(ws, [r"^t maxima hielo"]),
        "TVHF": _col_condicion_reg(ws, [r"maxima.*hielo.*viento"]),
        "TVM": _col_condicion_reg(ws, [r"viento 1 2"]),
        "15V": _col_condicion_reg(ws, [r"^15"]),
        "T0H": _col_condicion_reg(ws, [r"^0"]),
        "TTemp": _col_condicion_reg(ws, [r"^50"]),
    }
    for i, item_tramo in datos["tensiones"].items():
        fila = _fila_tramo_en(ws, col_tramo, i + 1, i + 2)
        if fila is None:
            continue
        for clave, t in item_tramo.items():
            col_t = cols_cond.get(clave)
            if col_t is None:
                continue
            ws.cell(row=fila, column=col_t + 1).value = round(t["T"], 1)
            if _tiene_flecha(ws, col_t):
                ws.cell(row=fila, column=col_t + 2).value = round(t["F"], 3)


def _ar_por_tramo(ruta_excel, vanos):
    """Vano de regulacion de cada tramo (el de su grupo, como mover_apoyo)."""
    n = len(vanos)
    grupos_clave = mv.leer_grupos_vanos_regulacion(ruta_excel)
    grupos_tramos = {}
    for i in range(n):
        g = grupos_clave.get((i + 1, i + 2), i + 1)
        grupos_tramos.setdefault(g, []).append(i)
    ar = {}
    for i in range(n):
        g = next(g for g, idxs in grupos_tramos.items() if i in idxs)
        valor = tv.vano_regulacion([vanos[j] for j in grupos_tramos[g]])
        ar[i] = valor if valor is not None else vanos[i]
    return ar


def escribir_tabla_temperaturas(ruta_excel, wb, datos, conductor):
    """Escribe T/F por temperatura en 'T. tend. cond. fase' (cabeceras de fila
    2 con un numero; T en fila 3 y F en la columna siguiente)."""
    ws = wb["T. tend. cond. fase"]
    col_tramo = _localizar_columna(ws, [r"^tramo$"])
    if col_tramo is None:
        return

    col_t_por_temp = {}
    for c in ws[2]:
        if c.value is None:
            continue
        if ws.cell(row=3, column=c.column).value is None:
            continue
        m = re.search(r"[-+]?\d+(?:[.,]\d+)?", str(c.value))
        if m:
            temp = float(m.group(0).replace(",", "."))
            col_t_por_temp[temp] = c.column - 1
    if not col_t_por_temp:
        return

    props = tv.propiedades_conductor(conductor)
    alfa = props["alfa_1_C"]
    E = props["E_daN_mm2"]
    S = props["seccion_mm2"]
    t0 = datos["t0_C"]
    vanos = datos["vanos"]
    desniveles = datos["desniveles"]
    p_cable = conductor["peso_daN_m"]
    ar_por_tramo = _ar_por_tramo(ruta_excel, vanos)

    for i in range(datos["n"]):
        fila = _fila_tramo_en(ws, col_tramo, i + 1, i + 2)
        if fila is None:
            continue
        T0 = datos["t0_por_tramo"][i] if i < len(datos["t0_por_tramo"]) else None
        if T0 is None:
            continue
        ar = ar_por_tramo[i]
        a = vanos[i]
        h = desniveles[i] if i < len(desniveles) else 0.0
        for temp, col_t in sorted(col_t_por_temp.items()):
            Tf = tv.tension_cambio_estado(t0, T0, p_cable, ar, temp, p_cable,
                                          alfa, E, S)
            if Tf is None:
                continue
            ws.cell(row=fila, column=col_t + 1).value = round(Tf, 1)
            if _tiene_flecha(ws, col_t):
                f = tv.flecha_vano(a, h, p_cable, Tf)
                ws.cell(row=fila, column=col_t + 2).value = round(f, 3)


# ---------------------------------------------------------------------------
# Cadenas: apoyos_configurados.json + conteos de Mediciones
# ---------------------------------------------------------------------------
def _categoria_mediciones(cadena_info):
    """('amarre'|'suspension', 'vidrio'|'polimericas') para las filas de
    Mediciones. SUS -> suspension, ANC -> amarre."""
    tipo = cadena_info.get("tipo", "")
    material = cadena_info.get("material", "")
    grupo = "amarre" if tipo == "ANC" else "suspension"
    tipo_mat = "vidrio" if material == "VID" else "polimericas"
    return grupo, tipo_mat


def leer_tipos_apoyos(wb):
    """Por apoyo de 'Cálculo de apoyos': tipo_clave, texto tipo, coef de
    seguridad, angulo, y coeficientes L, N, S. Devuelve {numero: dict}."""
    ws = wb["Cálculo de apoyos"]
    col_tipo = _localizar_columna(ws, [r"^tipo$"])
    col_ang = _localizar_columna(ws, [r"valor.*angulo"])
    col_seg = _localizar_columna(ws, [r"coefici.*seguri"])
    col_l = _localizar_columna(ws, [r"^coeficientes.*l.*n.*s"])
    col_n = col_l + 1 if col_l is not None else None
    col_s = col_n + 1 if col_n is not None else None
    resultado = {}
    for fila in ws.iter_rows(min_row=1):
        numero = _numero(fila[0].value)
        if numero is None:
            continue
        n = int(numero)
        tipo_texto = (_texto_limpio(fila[col_tipo].value)
                      if col_tipo is not None else None)
        angulo = _numero(fila[col_ang].value) if col_ang is not None else None
        coef_seg = _texto_limpio(fila[col_seg].value) if col_seg is not None else None
        l = _numero(fila[col_l].value) if col_l is not None else None
        nn = _numero(fila[col_n].value) if col_n is not None else None
        ss = _numero(fila[col_s].value) if col_s is not None else None
        resultado[n] = {
            "tipo_texto": tipo_texto,
            "tipo_clave": mapear_tipo_a_clave(tipo_texto),
            "angulo": angulo,
            "coef_seguridad": (coef_seg or "N").strip().upper()[:1],
            "L": l, "N": nn, "S": ss,
        }
    return resultado


def actualizar_apoyos_config(ruta_config, apoyos_afectados, cadena_info,
                             tipos_apoyos, ruta_salida=None):
    """Anade el bloque 'cadena' a los apoyos afectados de apoyos_configurados.json.
    'apoyos_afectados': lista de numeros de apoyo o "todos"."""
    if not os.path.exists(ruta_config):
        raise FileNotFoundError("No se encontro %s" % ruta_config)
    with open(ruta_config, encoding="utf-8-sig") as archivo:
        documento = json.load(archivo)

    car = cadena_info.get("caracteristicas", {})
    base_cadena = {
        "nombre": cadena_info["nombre"],
        "denominacion": cadena_info.get("denominacion", ""),
        "conductor": cadena_info.get("conductor"),
        "tension_kv": cadena_info.get("tension_kv"),
        "tipo": clasificar_cadena(cadena_info["nombre"])["tipo"],
        "material": clasificar_cadena(cadena_info["nombre"])["material"],
        "doble": clasificar_cadena(cadena_info["nombre"])["doble"],
        "caracteristicas": {
            "longitud_m": car.get("longitud_m"),
            "peso_daN": car.get("peso_daN"),
            "carga_rotura_daN": car.get("carga_rotura_daN"),
            "esfuerzo_viento_120_daN": car.get("esfuerzo_viento_120_daN"),
            "esfuerzo_viento_140_daN": car.get("esfuerzo_viento_140_daN"),
            "esfuerzo_viento_60_daN": car.get("esfuerzo_viento_60_daN"),
        },
    }

    modificados = []
    for apoyo in documento.get("apoyos", []):
        numero = (apoyo.get("apoyo") or {}).get("numero")
        if numero is None:
            continue
        if apoyos_afectados != "todos" and numero not in apoyos_afectados:
            continue
        tipo = tipos_apoyos.get(numero, {})
        ncad = n_cadenas(tipo.get("tipo_clave"), tipo.get("coef_seguridad", "N"))
        bloque = dict(base_cadena)
        bloque["n_cadenas"] = ncad
        apoyo["cadena"] = bloque
        modificados.append(numero)

    ruta_salida = ruta_salida or ruta_config
    with open(ruta_salida, "w", encoding="utf-8") as archivo:
        json.dump(documento, archivo, ensure_ascii=False, indent=2)
        archivo.write("\n")
    return modificados


def escribir_mediciones_cadenas(wb, apoyos):
    """Escribe en Mediciones el numero de cadenas por categoria (amarre o
    suspension x vidrio o polimericas). 'apoyos' trae el bloque 'cadena'."""
    ws = wb["Mediciones"]
    filas = {}
    for fila in ws.iter_rows():
        for c in fila:
            clave = _clave_normalizada(c.value)
            if not clave:
                continue
            if "cadenas de amarre de vidrio" in clave:
                filas[("amarre", "vidrio")] = fila
                break
            elif "cadenas de amarre polimericas" in clave:
                filas[("amarre", "polimericas")] = fila
                break
            elif "cadenas de suspension de vidrio" in clave:
                filas[("suspension", "vidrio")] = fila
                break
            elif "cadenas de suspension polimericas" in clave:
                filas[("suspension", "polimericas")] = fila
                break

    from collections import defaultdict
    conteos = defaultdict(int)
    for a in apoyos:
        cad = a.get("cadena") or {}
        if not cad:
            continue
        grupo, tipo_mat = _categoria_mediciones(cad)
        conteos[(grupo, tipo_mat)] += (cad.get("n_cadenas") or 1) * NUMERO_FASES

    for (grupo, tipo_mat) in (("amarre", "vidrio"), ("amarre", "polimericas"),
                              ("suspension", "vidrio"), ("suspension", "polimericas")):
        fila = filas.get((grupo, tipo_mat))
        if fila is None:
            continue
        destino = None
        for c in fila:
            if _texto_limpio(c.value) is not None:
                destino = c
        if destino is not None and destino.column > 1:
            destino.value = int(conteos.get((grupo, tipo_mat), 0))


# ---------------------------------------------------------------------------
# Esfuerzos de las hipotesis (1a-4a), desviacion de cadena y flecha maxima en
# 'Cálculo de apoyos'. Ecuaciones del Tema 10 (graficos_utilizacion_ecuaciones).
# ---------------------------------------------------------------------------
TIPOS_SUSPENSION = ("alineacion_suspension", "angulo_suspension")


def _tension_de_tramo(tensiones, i, clave):
    """T (daN) de la hipotesis 'clave' en el tramo de indice i (o None)."""
    if i is None or i < 0:
        return None
    item = tensiones.get(i, {})
    par = item.get(clave)
    return par["T"] if par else None


def _montar_variables_apoyo(k, tipos, vanos, desniveles, tensiones, sobre,
                            conductor, cadena, ncad):
    """Variables de las ecuaciones V/T/L del CalculoApoyos (HP Prime) para el
    apoyo k: vanos, pendientes, sobrecargas, tensiones por hipotesis, cadena,
    zona, clase de tension de linea y conductores por fase."""
    n = len(tipos)
    t = tipos[k]
    a1 = vanos[k - 2] if k >= 2 else None      # tramo anterior (k-1,k)
    a2 = vanos[k - 1] if k <= n - 1 else None  # tramo posterior (k,k+1)
    h1 = desniveles[k - 2] if k >= 2 else None
    h2 = desniveles[k - 1] if k <= n - 1 else None
    i1 = k - 2
    i2 = k - 1

    def _pendiente(a, h):
        return (h / a) if (a and a > 0) else 0.0

    def _cv(ik):
        if ik is None or ik < 0 or ik >= len(sobre):
            return None
        T = _tension_de_tramo(tensiones, ik, "TVF")
        sv = sobre[ik].get("Sv") or 0.0
        return (T / sv) if (T and sv) else None

    def _cvhf(ik):
        if ik is None or ik < 0 or ik >= len(sobre):
            return None
        T = _tension_de_tramo(tensiones, ik, "TVHF")
        svh = sobre[ik].get("S_vh") or 0.0
        return (T / svh) if (T and svh) else None

    def _rep(clave):
        """Valor representativo (media de los tramos adyacentes existentes)."""
        vals = []
        for ik in (i1, i2):
            if ik is not None and 0 <= ik < len(sobre):
                val = sobre[ik].get(clave)
                if val is not None:
                    vals.append(val)
        return (sum(vals) / len(vals)) if vals else None

    def _zona(ik):
        if ik is None or ik < 0 or ik >= len(sobre):
            return "A"
        return sobre[ik].get("zona", "A")

    car = (cadena or {}).get("caracteristicas", {})
    tension_kv = int((cadena or {}).get("tension_kv") or 20)
    zonas_lado = (_zona(i1), _zona(i2))
    return {
        "k": k,
        "tipo_clave": t["tipo_clave"],
        "angulo": t.get("angulo"),
        "a1": a1, "a2": a2, "n1": _pendiente(a1, h1), "n2": _pendiente(a2, h2),
        "i1": i1, "i2": i2,
        "CVF1": _cv(i1), "CVF2": _cv(i2),
        "CVHF1": _cvhf(i1), "CVHF2": _cvhf(i2),
        "TVF1": _tension_de_tramo(tensiones, i1, "TVF"),
        "TVF2": _tension_de_tramo(tensiones, i2, "TVF"),
        "THF1": _tension_de_tramo(tensiones, i1, "THF"),
        "THF2": _tension_de_tramo(tensiones, i2, "THF"),
        "TVHF1": _tension_de_tramo(tensiones, i1, "TVHF"),
        "TVHF2": _tension_de_tramo(tensiones, i2, "TVHF"),
        "TVM": max([x or 0.0 for x in
                    (_tension_de_tramo(tensiones, i1, "TVM"),
                     _tension_de_tramo(tensiones, i2, "TVM"))]) or None,
        "pF": conductor["peso_daN_m"],
        "dF": conductor["diametro_m"],
        "Sv": _rep("Sv"),
        "Sh": _rep("Sh"),
        "S_vh": _rep("S_vh"),
        "dmh": _rep("dmh"),
        "vF": _rep("P_v_max"),
        "vF60": _rep("V_vh"),
        "zona": ("C" if "C" in zonas_lado else
                 "B" if "B" in zonas_lado else "A"),
        "z1_zona": _zona(i1), "z2_zona": _zona(i2),
        "PCADF": car.get("peso_daN"),
        "EVCADF": car.get("esfuerzo_viento_120_daN"),
        "EVCADF60": car.get("esfuerzo_viento_60_daN"),
        "NCAD": ncad,
        "v_cat": 1 if tension_kv <= 66 else 2,
        "n_fase": CONDUCTORES_POR_FASE,
    }


def _raiz(*cuadrados):
    return math.sqrt(sum(c * c for c in cuadrados if c is not None))


# Conductores por fase (coeficiente de la 4a hipotesis de rotura):
# 1 = circuito simple (habitual); DOBLE_CIRCUITO = 2.
CONDUCTORES_POR_FASE = 1


def _es_angulo(v):
    return v["tipo_clave"] in ("angulo_anclaje", "angulo_amarre",
                               "angulo_suspension")


def _es_suspension(v):
    return v["tipo_clave"] in ("alineacion_suspension", "angulo_suspension")


def _es_anclaje(v):
    return v["tipo_clave"] in ("alineacion_anclaje", "angulo_anclaje")


def _es_pl(v):
    return v["tipo_clave"] == "principio_final_de_linea"


def _medio(v):
    return ((v["a1"] or 0.0) + (v["a2"] or 0.0)) / 2.0


def _pcad(v):
    return (v["PCADF"] or 0.0) * (v["NCAD"] or 0)


def _porcentaje_desequilibrio(tipo_clave, v_cat):
    """L_p3 (3a hipotesis): suspension 8/15 %, amarre 15/25 % (tension de la
    linea <=66 / >66 kV), anclaje siempre 50 % (CalculoApoyos)."""
    if tipo_clave in ("alineacion_anclaje", "angulo_anclaje"):
        return 0.50
    if tipo_clave in ("alineacion_amarre", "angulo_amarre"):
        return 0.15 if v_cat <= 1 else 0.25
    return 0.08 if v_cat <= 1 else 0.15


def _coef_rotura(tipo_clave, n_fase):
    """L_p4 (4a hipotesis): suspension 50/75/100 % (<=2/3/>=4 conductores),
    amarre 100 %, anclaje y P./F. linea 100 % (1) / 50 % (>=2)."""
    if tipo_clave in ("alineacion_suspension", "angulo_suspension"):
        if n_fase >= 4:
            return 1.00
        return 0.75 if n_fase == 3 else 0.50
    if tipo_clave in ("alineacion_amarre", "angulo_amarre"):
        return 1.00
    return 1.00 if n_fase == 1 else 0.50


def _traccion_angulo(v, T1, T2, es_suspension):
    """Componente de traccion de los apoyos de angulo (CalculoApoyos):
    suspension -> 2*max*cos(alfa/2); amarre/anclaje -> resultante vectorial."""
    alfa = v["angulo"] if v["angulo"] else 180.0
    c = math.cos(math.radians(alfa / 2.0))
    if es_suspension:
        return 2.0 * max(T1 or 0.0, T2 or 0.0) * c
    s = math.sin(math.radians(alfa / 2.0))
    return _raiz(((T1 or 0.0) + (T2 or 0.0)) * c,
                 ((T1 or 0.0) - (T2 or 0.0)) * s)


def _peso_tension_lado(v, lado, clave_th, clave_tv):
    """Peso (Sh o pF) y tension (THF o TVF) del lado indicado segun su zona."""
    zona = v["z1_zona"] if lado == 1 else v["z2_zona"]
    if zona in ("B", "C"):
        return v["Sh"], v[clave_th]
    return v["pF"], v[clave_tv]


def _esfuerzos_1a_viento(v):
    """1a hipotesis (viento): V, T, L para TODOS los tipos (CalculoApoyos)."""
    medio = _medio(v)
    V = v["pF"] * (medio + (v["CVF1"] or 0.0) * (v["n1"] or 0.0)
                   - (v["CVF2"] or 0.0) * (v["n2"] or 0.0)) + _pcad(v)
    if _es_angulo(v):
        alfa = v["angulo"] if v["angulo"] else 180.0
        c_beta = math.cos(math.radians((180.0 - alfa) / 2.0))
        T = v["dF"] * (v["vF"] or 0.0) * medio * c_beta \
            + (v["EVCADF"] or 0.0) * (v["NCAD"] or 0) \
            + _traccion_angulo(v, v["TVF1"], v["TVF2"], _es_suspension(v))
    else:
        T = v["dF"] * (v["vF"] or 0.0) * medio \
            + (v["EVCADF"] or 0.0) * (v["NCAD"] or 0)
    L = max(v["TVF1"] or 0.0, v["TVF2"] or 0.0) if _es_pl(v) else None
    return V, T, L


def _esfuerzos_2a_hielo(v):
    """2a hipotesis (hielo, zonas B/C): V2h, T2h (angulo) y L2h (P./F. linea)."""
    if v["zona"] in (None, "A"):
        return None, None, None
    a1 = v["a1"] or 0.0
    a2 = v["a2"] or 0.0
    W1, T1 = _peso_tension_lado(v, 1, "THF1", "TVF1")
    W2, T2 = _peso_tension_lado(v, 2, "THF2", "TVF2")
    V = (W1 or 0.0) * (a1 / 2.0) + (T1 or 0.0) * (v["n1"] or 0.0) \
        + (W2 or 0.0) * (a2 / 2.0) - (T2 or 0.0) * (v["n2"] or 0.0) + _pcad(v)
    T = None
    if _es_angulo(v):
        T = _traccion_angulo(v, T1, T2, _es_suspension(v))
    L = max(v["THF1"] or 0.0, v["THF2"] or 0.0) if _es_pl(v) else None
    return V, T, L


def _esfuerzos_2a_vh(v):
    """2a hipotesis viento+hielo (zonas B/C). Solo si existe TVHF en la linea."""
    if v["zona"] in (None, "A"):
        return None, None, None
    if v["TVHF1"] is None and v["TVHF2"] is None:
        return None, None, None
    medio = _medio(v)
    if _es_anclaje(v) and not _es_angulo(v):
        # Anclaje de alineacion: formula por lado (sin dividir por S_vh).
        W1, T1 = _peso_tension_lado(v, 1, "TVHF1", "TVF1")
        W2, T2 = _peso_tension_lado(v, 2, "TVHF2", "TVF2")
        V = (W1 or 0.0) * ((v["a1"] or 0.0) / 2.0) + (T1 or 0.0) * (v["n1"] or 0.0) \
            + (W2 or 0.0) * ((v["a2"] or 0.0) / 2.0) - (T2 or 0.0) * (v["n2"] or 0.0) \
            + _pcad(v)
    else:
        V = (v["Sh"] or 0.0) * (medio + (v["CVHF1"] or 0.0) * (v["n1"] or 0.0)
                                - (v["CVHF2"] or 0.0) * (v["n2"] or 0.0)) + _pcad(v)
    base = (v["dmh"] or v["dF"] or 0.0) * (v["vF60"] or 0.0) * medio \
        + (v["EVCADF60"] or 0.0) * (v["NCAD"] or 0)
    if _es_angulo(v):
        alfa = v["angulo"] if v["angulo"] else 180.0
        c_beta = math.cos(math.radians((180.0 - alfa) / 2.0))
        T = (v["dmh"] or v["dF"] or 0.0) * (v["vF60"] or 0.0) * medio * c_beta \
            + (v["EVCADF60"] or 0.0) * (v["NCAD"] or 0) \
            + _traccion_angulo(v, v["TVHF1"], v["TVHF2"], _es_suspension(v))
    else:
        T = base
    L = max(v["TVHF1"] or 0.0, v["TVHF2"] or 0.0) if _es_pl(v) else None
    return V, T, L


def _esfuerzos_3a_desequilibrio(v):
    """3a hipotesis (desequilibrio). No aplica a P./F. de linea."""
    if _es_pl(v):
        return None, None, None
    p = _porcentaje_desequilibrio(v["tipo_clave"], v["v_cat"])
    if v["zona"] in (None, "A"):
        V3 = v["pF"] * (_medio(v) + (v["CVF1"] or 0.0) * (v["n1"] or 0.0)
                        - (v["CVF2"] or 0.0) * (v["n2"] or 0.0)) + _pcad(v)
        T1, T2 = v["TVF1"], v["TVF2"]
        T3 = (_traccion_angulo(v, T1, T2, _es_suspension(v))
              if _es_angulo(v) else None)
        L3 = p * max(T1 or 0.0, T2 or 0.0)
    else:
        # Opcion Hielo (la que recoge el Excel para zonas B/C).
        V3, T3, _ = _esfuerzos_2a_hielo(v)
        T1, T2 = v["THF1"], v["THF2"]
        L3 = p * max(T1 or 0.0, T2 or 0.0)
    return V3, T3, L3


def _esfuerzos_4a_rotura(v):
    """4a hipotesis (rotura): devuelve V(NA), V(A), T(NA), T(A), LFA, LFNA."""
    p4 = _coef_rotura(v["tipo_clave"], v["n_fase"])
    if v["zona"] in (None, "A"):
        VNA, TNA, _ = _esfuerzos_3a_desequilibrio(v)
        T1, T2 = v["TVF1"], v["TVF2"]
    else:
        VNA, TNA, _ = _esfuerzos_2a_hielo(v)
        T1, T2 = v["THF1"], v["THF2"]
    VA = 0.5 * VNA if VNA is not None else None
    TA = 0.5 * TNA if TNA is not None else None
    if _es_pl(v):
        LFNA = p4 * max(T1 or 0.0, T2 or 0.0)
        LFA = 0.0
    else:
        LFNA = 0.0
        LFA = p4 * max(T1 or 0.0, T2 or 0.0)
        if v["tipo_clave"] in ("angulo_anclaje", "angulo_amarre"):
            alfa = v["angulo"] if v["angulo"] else 180.0
            LFA *= math.sin(math.radians(alfa / 2.0))
    return VNA, VA, TNA, TA, LFA, LFNA


def _columnas_esfuerzos(ws):
    """Columnas (0-based) de los bloques de 'Cálculo de apoyos'."""
    def _col(patrones, fila=1):
        return _localizar_columna(ws, patrones, fila_cabecera=fila)

    bloques = {
        "1a": _col([r"^1a hipotesis"]),
        "2a_hielo": None,
        "2a_vh": None,
        "3a": _col([r"^3a hipotesis"]),
        "4a_no": _col([r"fases no afectadas"], fila=2),
        "4a_si": _col([r"fases afectadas"], fila=2),
        "esf_tor": _col([r"esf.*tor"], fila=2),
        "aa": _col([r"desviaci.*cadena"]),
        "ab": _col([r"flecha.*maxima"]),
    }
    for c in ws[2]:
        clave = _clave_normalizada(c.value)
        if not clave:
            continue
        if clave == "hielo":
            bloques["2a_hielo"] = c.column - 1
        elif "hielo" in clave and "viento" in clave:
            bloques["2a_vh"] = c.column - 1
    return bloques


def _redondeo(x):
    """Redondeo comercial a entero (como muestra el Excel del TFG)."""
    return math.floor(x + 0.5) if x is not None else None


def _poner(ws, fila, col_bloque, V, T, L):
    """Escribe la terna V/T/L a partir de la columna V del bloque."""
    if col_bloque is None:
        return
    ws.cell(row=fila, column=col_bloque + 1).value = _redondeo(V)
    ws.cell(row=fila, column=col_bloque + 2).value = _redondeo(T)
    ws.cell(row=fila, column=col_bloque + 3).value = _redondeo(L)


def _fila_de_apoyo(ws, k):
    for r in ws.iter_rows(min_row=1, max_col=2):
        if _numero(r[0].value) == k:
            return r[0].row
    return None


def escribir_esfuerzos_apoyos(wb, datos, tipos, cadenas_por_apoyo):
    """Reescribe los esfuerzos V/T/L de las hipotesis 1a-4a (y el esfuerzo
    torsor aplicado) en 'Cálculo de apoyos' segun CalculoApoyos (HP Prime).
    'cadenas_por_apoyo': {numero: bloque cadena}."""
    ws = wb["Cálculo de apoyos"]
    cols = _columnas_esfuerzos(ws)
    if cols["1a"] is None:
        return

    por_apoyo = {}
    for k in sorted(tipos.keys()):
        cadena = cadenas_por_apoyo.get(k)
        ncad = (cadena or {}).get("n_cadenas", 1)
        por_apoyo[k] = _montar_variables_apoyo(
            k, tipos, datos["vanos"], datos["desniveles"], datos["tensiones"],
            datos["sobrecargas"], datos["conductor"], cadena, ncad)

    for k, v in por_apoyo.items():
        fila = _fila_de_apoyo(ws, k)
        if fila is None:
            continue

        V1, T1, L1 = _esfuerzos_1a_viento(v)
        _poner(ws, fila, cols["1a"], V1, T1, L1 or None)

        if cols["2a_hielo"] is not None:
            V2, T2, L2 = _esfuerzos_2a_hielo(v)
            _poner(ws, fila, cols["2a_hielo"], V2, T2, L2 or None)

        if cols["2a_vh"] is not None:
            V2v, T2v, L2v = _esfuerzos_2a_vh(v)
            _poner(ws, fila, cols["2a_vh"], V2v, T2v, L2v or None)

        if cols["3a"] is not None:
            V3, T3, L3 = _esfuerzos_3a_desequilibrio(v)
            _poner(ws, fila, cols["3a"], V3, T3, L3 or None)

        VNA, VA, TNA, TA, LFA, LFNA = _esfuerzos_4a_rotura(v)
        if cols["4a_no"] is not None:
            _poner(ws, fila, cols["4a_no"], VNA, TNA, LFNA or None)
        if cols["4a_si"] is not None:
            _poner(ws, fila, cols["4a_si"], VA, TA, LFA or None)
        if cols["esf_tor"] is not None:
            torsor = max(LFNA or 0.0, LFA or 0.0)
            ws.cell(row=fila, column=cols["esf_tor"] + 1).value = (
                _redondeo(torsor) if torsor else None)

    return por_apoyo


def _formato_par(a, b):
    """Formato 'afectada/no afectada' de la 4a hipotesis en 'Elección apoyos'."""
    sa = str(_redondeo(a)) if a is not None else "——"
    sb = str(_redondeo(b)) if b is not None else "——"
    return "%s/%s" % (sa, sb)


def escribir_eleccion_apoyos_esfuerzos(wb, tipos, datos, cadenas_por_apoyo):
    """Actualiza el bloque 'Esfuerzo por fase y tierra' de 'Elección apoyos'
    (V/T/L de las filas 'Fase' de cada hipotesis) con los mismos valores de
    'Cálculo de apoyos' (CalculoApoyos)."""
    if "Elección apoyos" not in wb.sheetnames:
        return
    ws = wb["Elección apoyos"]
    por_apoyo = {}
    for k in sorted(tipos.keys()):
        cadena = cadenas_por_apoyo.get(k)
        ncad = (cadena or {}).get("n_cadenas", 1)
        por_apoyo[k] = _montar_variables_apoyo(
            k, tipos, datos["vanos"], datos["desniveles"], datos["tensiones"],
            datos["sobrecargas"], datos["conductor"], cadena, ncad)

    max_fila = ws.max_row or 1
    fila = 1
    while fila <= max_fila:
        numero = _numero(ws.cell(row=fila, column=1).value)
        if numero is None or numero not in por_apoyo:
            fila += 1
            continue
        v = por_apoyo[numero]
        f = fila
        while f <= max_fila:
            if f > fila and _numero(ws.cell(row=f, column=1).value) is not None:
                break
            etiqueta = _clave_normalizada(ws.cell(row=f, column=8).value)
            conductor = _clave_normalizada(ws.cell(row=f, column=9).value)
            if conductor != "fase" or not etiqueta:
                f += 1
                continue
            if "rotu" in etiqueta:
                VNA, VA, TNA, TA, LFA, LFNA = _esfuerzos_4a_rotura(v)
                ws.cell(row=f, column=10).value = _formato_par(VA, VNA)
                ws.cell(row=f, column=11).value = _formato_par(TA, TNA)
                ws.cell(row=f, column=12).value = _redondeo(
                    max(LFNA or 0.0, LFA or 0.0) or None)
            elif "dese" in etiqueta:
                V, T, L = _esfuerzos_3a_desequilibrio(v)
                ws.cell(row=f, column=10).value = _redondeo(V)
                ws.cell(row=f, column=11).value = _redondeo(T)
                ws.cell(row=f, column=12).value = _redondeo(L or None)
            elif "vien" in etiqueta and "hielo" in etiqueta:
                V, T, L = _esfuerzos_2a_vh(v)
                ws.cell(row=f, column=10).value = _redondeo(V)
                ws.cell(row=f, column=11).value = _redondeo(T)
                ws.cell(row=f, column=12).value = _redondeo(L or None)
            elif "hielo" in etiqueta:
                V, T, L = _esfuerzos_2a_hielo(v)
                ws.cell(row=f, column=10).value = _redondeo(V)
                ws.cell(row=f, column=11).value = _redondeo(T)
                ws.cell(row=f, column=12).value = _redondeo(L or None)
            elif "vien" in etiqueta:
                V, T, L = _esfuerzos_1a_viento(v)
                ws.cell(row=f, column=10).value = _redondeo(V)
                ws.cell(row=f, column=11).value = _redondeo(T)
                ws.cell(row=f, column=12).value = _redondeo(L or None)
            f += 1
        fila = f


def recalcular_esfuerzos_apoyos(ruta_excel, salida_excel=None, tensiones=None,
                                conductor=None, viento_kmh=VIENTO_KMH):
    """Recalcula (o reutiliza) las tensiones y reescribe los esfuerzos V/T/L de
    las hipotesis 1a-4a en 'Cálculo de apoyos' y el esfuerzo por fase de
    'Elección apoyos', segun CalculoApoyos (HP Prime). Devuelve la ruta.
    'tensiones': opcional, {i_tramo: {clave: {'T','F'}}} (p.ej. las que acaba
    de recalcular mover_apoyo o modulos_fuste tras cambiar la geometria)."""
    salida = salida_excel or ruta_excel
    if salida != ruta_excel and not os.path.exists(salida):
        shutil.copy2(ruta_excel, salida)
    wb = openpyxl.load_workbook(salida)
    try:
        if conductor is None:
            conductor = _leer_conductor_actual(ruta_excel)
        if tensiones is None:
            datos = recalcular_tensiones_linea(ruta_excel, conductor, viento_kmh)
        else:
            geom = mv.leer_geometria_excel(ruta_excel)
            datos = {
                "tensiones": tensiones,
                "vanos": [t[0] for t in geom["tramos"]],
                "desniveles": [t[1] for t in geom["tramos"]],
                "n": geom["n"],
                "sobrecargas": sobrecargas_por_tramo(ruta_excel, conductor,
                                                     viento_kmh),
            }
        datos["conductor"] = conductor
        tipos = leer_tipos_apoyos(wb)
        cadenas = _cadenas_por_apoyo(RUTA_APOYOS_CONFIG, tipos, conductor,
                                     conductor["nomenclatura_excel"])
        escribir_esfuerzos_apoyos(wb, datos, tipos, cadenas)
        escribir_desviacion_cadena(wb, datos, tipos, cadenas)
        escribir_flecha_maxima(wb, datos)
        escribir_eleccion_apoyos_esfuerzos(wb, tipos, datos, cadenas)
        wb.save(salida)
        regenerar_graficos_utilizacion(salida, cadenas)
    finally:
        wb.close()
    return os.path.abspath(salida)


def regenerar_graficos_utilizacion(ruta_excel, cadenas_por_apoyo=None):
    """Regenera graficos_utilizacion.json y los PNG de la ventana Unity con las
    cadenas ACTUALES del gemelo (apoyos_configurados.json). No rompe el flujo
    si falta algun modulo o dato."""
    try:
        from generar_graficos_utilizacion import generar_json
        from generar_png_graficos import generar_pngs
        nombres = None
        if cadenas_por_apoyo:
            nombres = {k: (c or {}).get("nombre")
                       for k, c in cadenas_por_apoyo.items()
                       if (c or {}).get("nombre")}
        generar_json(ruta_excel, cadenas_por_apoyo=nombres)
        generar_pngs()
        print("  [gemelo] Graficos de utilizacion regenerados.")
        return True
    except Exception as error_graf:
        print("  ⚠️ No se pudieron regenerar los graficos de utilizacion:",
              error_graf)
        return False

def escribir_desviacion_cadena(wb, datos, tipos, cadenas_por_apoyo):
    """Desviacion de la cadena (col. AA) para apoyos de suspension:
       tan(gamma) = [v/2*d*L + EVCad/2] / [p*L + TVM*N + PCad/2]"""
    ws = wb["Cálculo de apoyos"]
    col_aa = _localizar_columna(ws, [r"desviaci.*cadena"])
    if col_aa is None:
        return
    for k in sorted(tipos.keys()):
        tipo_clave = tipos[k]["tipo_clave"]
        if tipo_clave not in TIPOS_SUSPENSION:
            continue
        cadena = cadenas_por_apoyo.get(k) or {}
        car = cadena.get("caracteristicas", {})
        v = _montar_variables_apoyo(
            k, tipos, datos["vanos"], datos["desniveles"], datos["tensiones"],
            datos["sobrecargas"], datos["conductor"], cadena,
            (cadena.get("n_cadenas") or 1))
        medio = ((v["a1"] or 0.0) + (v["a2"] or 0.0)) / 2.0
        pcad = car.get("peso_daN")
        evcad = car.get("esfuerzo_viento_120_daN")
        if pcad is None or evcad is None or medio <= 0:
            continue
        numerador = (v["vF"] or 0.0) / 2.0 * v["dF"] * medio + evcad / 2.0
        denominador = (v["pF"] * medio
                       + (v["TVM"] or 0.0) * ((v["n1"] or 0.0) - (v["n2"] or 0.0))
                       + pcad / 2.0)
        if denominador <= 0:
            continue
        gamma = math.degrees(math.atan(numerador / denominador))
        fila = _fila_de_apoyo(ws, k)
        if fila is not None:
            ws.cell(row=fila, column=col_aa + 1).value = round(gamma, 2)


def escribir_flecha_maxima(wb, datos):
    """Flecha maxima por apoyo (col. AB): maximo de las flechas a 50 C de sus
    tramos adyacentes (hipotesis TTemp)."""
    ws = wb["Cálculo de apoyos"]
    col_ab = _localizar_columna(ws, [r"flecha.*maxima"])
    if col_ab is None:
        return
    n = datos["n"] + 1
    flecha_por_tramo = {}
    for i, item in datos["tensiones"].items():
        par = item.get("TTemp")
        flecha_por_tramo[i] = par["F"] if par else None
    for k in range(1, n + 1):
        valores = []
        if k >= 2 and (k - 2) in flecha_por_tramo:
            valores.append(flecha_por_tramo[k - 2])
        if k <= n - 1 and (k - 1) in flecha_por_tramo:
            valores.append(flecha_por_tramo[k - 1])
        validos = [v for v in valores if v is not None]
        if not validos:
            continue
        fila = _fila_de_apoyo(ws, k)
        if fila is not None:
            ws.cell(row=fila, column=col_ab + 1).value = round(max(validos), 3)


# ---------------------------------------------------------------------------
# conductor_configurado.json (Unity) - replica la logica de principal.py
# ---------------------------------------------------------------------------
def escribir_conductor_config(nomenclatura, ruta_salida=None):
    """Regenera conductor_configurado.json con el nuevo conductor."""
    coincidencia = resolver_conductor(nomenclatura)
    catalogo = cargar_catalogo_conductores()
    propiedades_composicion = {
        e["composicion"]: e
        for e in catalogo.get("propiedades_mecanicas_por_composicion", [])
    }
    props = propiedades_composicion.get(coincidencia["composicion"], coincidencia)

    documento = {
        "version_esquema": 3,
        "nomenclatura_excel": coincidencia["nomenclatura_excel"],
        "designacion": coincidencia["designacion"],
        "composicion": coincidencia["composicion"],
        "diametro_mm": coincidencia["diametro_mm"],
        "diametro_m": coincidencia["diametro_m"],
        "seccion_mm2": coincidencia["seccion_mm2"],
        "masa_kg_km": coincidencia["masa_kg_km"],
        "peso_daN_m": coincidencia["peso_daN_m"],
        "modulo_elasticidad_daN_mm2": props["modulo_elasticidad_daN_mm2"],
        "coef_dilatacion_1_C": props["coef_dilatacion_1_C"],
        "propiedades_mecanicas_verificadas": True,
        "fuente_propiedades_mecanicas": coincidencia["fuente_propiedades"],
        "numero_fases": NUMERO_FASES,
        "fuente_catalogo": os.path.abspath(RUTA_CATALOGO_CONDUCTORES),
    }
    # Enriquecer con resistencias electricas (DLR) si estan en la tabla
    try:
        if os.path.exists(RUTA_RESISTENCIAS):
            with open(RUTA_RESISTENCIAS, encoding="utf-8") as archivo:
                tabla = json.load(archivo).get("conductores", {})
            registro = tabla.get(_clave_conductor(nomenclatura))
            if registro is None:
                clave_desig = _clave_conductor(coincidencia["designacion"])
                for candidato in tabla.values():
                    if _clave_conductor(str(candidato.get("codigo", ""))) == clave_desig:
                        registro = candidato
                        break
            if registro:
                documento["r_low_ohm_km"] = registro["r_low_ohm_km"]
                documento["r_high_ohm_km"] = registro["r_high_ohm_km"]
                documento["t_low_C"] = registro.get("t_low_C", 25.0)
                documento["t_high_C"] = registro.get("t_high_C", 70.0)
                documento["fuente_resistencias"] = (
                    "data/catalogs/conductors/resistencias_electricas.json")
            else:
                documento["r_low_ohm_km"] = None
                documento["r_high_ohm_km"] = None
                documento["t_low_C"] = 25.0
                documento["t_high_C"] = 70.0
                print("  ⚠️ Sin resistencias electricas para '%s' (DLR sin R)."
                      % nomenclatura)
    except Exception as error:
        print("  ⚠️ No se pudo enriquecer el conductor con resistencias:", error)

    ruta_salida = ruta_salida or RUTA_CONDUCTOR_CONFIG
    with open(ruta_salida, "w", encoding="utf-8") as archivo:
        json.dump(documento, archivo, ensure_ascii=False, indent=2)
        archivo.write("\n")
    return documento


# ---------------------------------------------------------------------------
# Orquestador
# ---------------------------------------------------------------------------
def _cadena_por_defecto(conductor, tipo_clave):
    """Primera cadena del catalogo compatible con el conductor y el tipo de
    apoyo (SUS para suspension, ANC para el resto). Sirve de respaldo mientras
    un apoyo no tenga cadena asignada en apoyos_configurados.json."""
    sus = tipo_clave in TIPOS_SUSPENSION
    for c in cargar_catalogo_cadenas():
        if _clave_conductor(c.get("conductor")) != _clave_conductor(conductor):
            continue
        cl = clasificar_cadena(c["nombre"])
        if sus and cl["tipo"] == "SUS":
            return c
        if not sus and cl["tipo"] == "ANC":
            return c
    return None


def _leer_conductor_actual(ruta_excel):
    """Conductor actual del Excel, resuelto contra el catalogo."""
    from datos_apoyos_excel import leer_conductor
    leido = leer_conductor(ruta_excel)
    return resolver_conductor(leido["nomenclatura_excel"])


def _cadenas_por_apoyo(ruta_apoyos, tipos, conductor, nuevo_nombre):
    """Bloque cadena por apoyo: el asignado en apoyos_configurados.json o,
    si no existe, el primero del catalogo compatible (SUS/ANC segun el tipo)."""
    cadenas_por_apoyo = {}
    if not os.path.exists(ruta_apoyos):
        return cadenas_por_apoyo
    with open(ruta_apoyos, encoding="utf-8-sig") as archivo:
        doc_apoyos = json.load(archivo)
    for apoyo in doc_apoyos.get("apoyos", []):
        num = (apoyo.get("apoyo") or {}).get("numero")
        if num is None:
            continue
        cad = apoyo.get("cadena")
        if not cad:
            tipo_clave = tipos.get(num, {}).get("tipo_clave")
            por_defecto = _cadena_por_defecto(nuevo_nombre, tipo_clave)
            if por_defecto:
                car = por_defecto.get("caracteristicas", {})
                cad = {
                    "nombre": por_defecto["nombre"],
                    "n_cadenas": n_cadenas(
                        tipo_clave,
                        tipos.get(num, {}).get("coef_seguridad", "N")),
                    "caracteristicas": {
                        "peso_daN": car.get("peso_daN"),
                        "esfuerzo_viento_120_daN": car.get(
                            "esfuerzo_viento_120_daN"),
                        "esfuerzo_viento_60_daN": car.get(
                            "esfuerzo_viento_60_daN"),
                    },
                }
        cadenas_por_apoyo[num] = cad
    return cadenas_por_apoyo


def sincronizar_tipos_cadenas(ruta_excel, ruta_apoyos=None, ruta_conductor=None,
                              ruta_salida=None):
    """Sincroniza apoyos_configurados.json con el Excel de la linea:
      - Escribe 'tipo_clave', 'tipo_texto', 'coef_seguridad' y 'angulo' de
        cada apoyo (hoja 'Calculo de apoyos').
      - Rellena el bloque 'cadena' de los apoyos que aun no lo tienen con la
        primera cadena del catalogo compatible con el conductor vigente:
        SUS para apoyos de suspension, ANC (amarre/anclaje) para el resto.
    Es el paso previo para que Unity instale las cadenas (suspension vertical;
    amarre/anclaje siguiendo la direccion del cable)."""
    ruta_apoyos = ruta_apoyos or RUTA_APOYOS_CONFIG
    ruta_conductor = ruta_conductor or RUTA_CONDUCTOR_CONFIG
    if not os.path.exists(ruta_excel):
        raise FileNotFoundError("Excel de trabajo no encontrado: %s" % ruta_excel)
    if not os.path.exists(ruta_apoyos):
        raise FileNotFoundError("apoyos_configurados.json no encontrado: %s"
                                % ruta_apoyos)

    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    tipos = leer_tipos_apoyos(wb)
    wb.close()

    with open(ruta_apoyos, encoding="utf-8-sig") as archivo:
        documento = json.load(archivo)

    resumen = enriquecer_documento_apoyos(
        documento, ruta_excel, ruta_conductor=ruta_conductor)

    ruta_salida = ruta_salida or ruta_apoyos
    with open(ruta_salida, "w", encoding="utf-8") as archivo:
        json.dump(documento, archivo, ensure_ascii=False, indent=2)
        archivo.write("\n")

    resumen["ok"] = True
    resumen["apoyos_config"] = os.path.abspath(ruta_salida)
    return resumen


def enriquecer_documento_apoyos(documento, ruta_excel, ruta_conductor=None):
    """Anade a cada apoyo del documento apoyos_configurados.json los campos
    'tipo_clave', 'tipo_texto', 'coef_seguridad' y 'angulo' (hoja 'Calculo de
    apoyos') y rellena el bloque 'cadena' por defecto si no existe (SUS para
    suspension, ANC para el resto). Devuelve un dict resumen. Se llama desde
    sincronizar_tipos_cadenas y desde principal.py al regenerar el JSON."""
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    tipos = leer_tipos_apoyos(wb)
    wb.close()

    conductor = _leer_conductor_config(ruta_conductor, ruta_excel)

    tipos_aplicados = 0
    cadenas_asignadas = 0
    for apoyo in documento.get("apoyos", []):
        numero = (apoyo.get("apoyo") or {}).get("numero")
        if numero is None:
            continue
        tipo = tipos.get(numero, {})
        for campo in ("tipo_clave", "tipo_texto", "coef_seguridad", "angulo"):
            valor = tipo.get(campo)
            if valor is not None:
                apoyo[campo] = valor
        if tipo.get("tipo_clave"):
            tipos_aplicados += 1

        if not apoyo.get("cadena") and conductor:
            por_defecto = _cadena_por_defecto(conductor, tipo.get("tipo_clave"))
            if por_defecto:
                apoyo["cadena"] = _bloque_cadena(por_defecto, tipo)
                cadenas_asignadas += 1

    return {
        "apoyos_sincronizados": len(documento.get("apoyos", [])),
        "tipos_aplicados": tipos_aplicados,
        "cadenas_asignadas_por_defecto": cadenas_asignadas,
        "conductor": conductor,
    }


def _leer_conductor_config(ruta_conductor, ruta_excel):
    """Nomenclatura del conductor vigente: el de conductor_configurado.json o,
    si falta, el leido del Excel (resuelto contra el catalogo)."""
    if ruta_conductor and os.path.exists(ruta_conductor):
        try:
            with open(ruta_conductor, encoding="utf-8-sig") as archivo:
                doc = json.load(archivo)
            nomenclatura = doc.get("nomenclatura_excel") or doc.get("designacion")
            if nomenclatura:
                return nomenclatura
        except Exception:
            pass
    try:
        from datos_apoyos_excel import leer_conductor
        return leer_conductor(ruta_excel)["nomenclatura_excel"]
    except Exception:
        return None


def _bloque_cadena(cadena_info, tipo):
    """Bloque 'cadena' para apoyos_configurados.json (misma forma que
    actualizar_apoyos_config)."""
    car = cadena_info.get("caracteristicas", {})
    return {
        "nombre": cadena_info["nombre"],
        "denominacion": cadena_info.get("denominacion", ""),
        "conductor": cadena_info.get("conductor"),
        "tension_kv": cadena_info.get("tension_kv"),
        "tipo": clasificar_cadena(cadena_info["nombre"])["tipo"],
        "material": clasificar_cadena(cadena_info["nombre"])["material"],
        "doble": clasificar_cadena(cadena_info["nombre"])["doble"],
        "n_cadenas": n_cadenas(tipo.get("tipo_clave"),
                               tipo.get("coef_seguridad", "N")),
        "caracteristicas": {
            "longitud_m": car.get("longitud_m"),
            "peso_daN": car.get("peso_daN"),
            "carga_rotura_daN": car.get("carga_rotura_daN"),
            "esfuerzo_viento_120_daN": car.get("esfuerzo_viento_120_daN"),
            "esfuerzo_viento_140_daN": car.get("esfuerzo_viento_140_daN"),
            "esfuerzo_viento_60_daN": car.get("esfuerzo_viento_60_daN"),
        },
    }


def aplicar_peticion(p):
    """Peticion: {"excel":.., "salida_excel":.., "conductor": "LA-78",
    "cadena": {"nombre":.., "ambito": "apoyo|todos", "apoyo": N},
    "apoyos_config":.., "config_conductor":..}. Devuelve un dict resumen."""
    excel = p["excel"]
    salida = p.get("salida_excel") or excel
    ruta_apoyos = p.get("apoyos_config") or RUTA_APOYOS_CONFIG
    ruta_cond = p.get("config_conductor") or RUTA_CONDUCTOR_CONFIG

    if not os.path.exists(excel):
        raise FileNotFoundError("Excel de trabajo no encontrado: %s" % excel)
    if salida != excel and not os.path.exists(salida):
        shutil.copy2(excel, salida)

    wb = openpyxl.load_workbook(salida)
    resumen = {"ok": True, "salida_excel": os.path.abspath(salida)}

    # 1) Conductor nuevo (o el actual si solo cambian las cadenas)
    if p.get("conductor"):
        nuevo = resolver_conductor(p["conductor"])
        anterior = _leer_conductor_actual(excel)
        escribir_conductor_en_excel(wb, nuevo["nomenclatura_excel"])
        escribir_conductor_config(p["conductor"], ruta_cond)
        resumen["conductor"] = {
            "anterior": anterior["nomenclatura_excel"],
            "nuevo": nuevo["nomenclatura_excel"],
            "designacion": nuevo["designacion"],
        }
    else:
        nuevo = _leer_conductor_actual(salida)

    # 2) Tensiones y flechas con el conductor vigente (EDS constante)
    datos = recalcular_tensiones_linea(salida, nuevo)
    datos["conductor"] = nuevo
    escribir_tensiones_reg(wb, datos)
    escribir_tabla_temperaturas(salida, wb, datos, nuevo)

    # 3) Cadenas (apoyos_configurados.json)
    tipos = leer_tipos_apoyos(wb)
    if p.get("cadena"):
        cad = p["cadena"]
        cadena_info = buscar_cadena(cad["nombre"])
        if cadena_info is None:
            raise ValueError("Cadena '%s' no existe en el catalogo." % cad["nombre"])
        ambito = cad.get("ambito", "apoyo")
        afectados = "todos" if ambito == "todos" else [int(cad.get("apoyo", 1))]
        modificados = actualizar_apoyos_config(
            ruta_apoyos, afectados, cadena_info, tipos)
        resumen["cadena"] = {
            "nombre": cadena_info["nombre"],
            "ambito": ambito,
            "apoyos_modificados": modificados,
        }

    # 4) Cadenas por apoyo para los esfuerzos (asignadas o por defecto)
    cadenas_por_apoyo = _cadenas_por_apoyo(
        ruta_apoyos, tipos, nuevo, nuevo["nomenclatura_excel"])

    # 5) Esfuerzos, desviacion de cadena y flecha maxima
    escribir_esfuerzos_apoyos(wb, datos, tipos, cadenas_por_apoyo)
    escribir_desviacion_cadena(wb, datos, tipos, cadenas_por_apoyo)
    escribir_flecha_maxima(wb, datos)
    escribir_eleccion_apoyos_esfuerzos(wb, tipos, datos, cadenas_por_apoyo)

    # 6) Conteos de cadenas en Mediciones
    if p.get("cadena") and os.path.exists(ruta_apoyos):
        with open(ruta_apoyos, encoding="utf-8-sig") as archivo:
            doc_apoyos = json.load(archivo)
        escribir_mediciones_cadenas(wb, doc_apoyos.get("apoyos", []))

    wb.save(salida)
    regenerar_graficos_utilizacion(salida, cadenas_por_apoyo)
    resumen["apoyos_config"] = os.path.abspath(ruta_apoyos)
    resumen["config_conductor"] = os.path.abspath(ruta_cond)
    return resumen


# ---------------------------------------------------------------------------
# CLI (JSON) para Unity
# ---------------------------------------------------------------------------
def procesar_peticion_json(ruta_peticion, ruta_resultado=None):
    with open(ruta_peticion, encoding="utf-8-sig") as archivo:
        p = json.load(archivo)
    try:
        resultado = aplicar_peticion(p)
    except Exception as error:
        resultado = {"ok": False, "error": str(error)}
    if ruta_resultado:
        with open(ruta_resultado, "w", encoding="utf-8") as archivo:
            json.dump(resultado, archivo, ensure_ascii=False, indent=2,
                      default=str)
    return resultado


if __name__ == "__main__":
    import sys
    import io
    try:
        if sys.stdout.encoding and "utf" not in sys.stdout.encoding.lower():
            sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                          errors="replace")
    except Exception:
        pass

    if "--catalogo_cadenas" in sys.argv:
        filtro = None
        if "--conductor" in sys.argv:
            filtro = sys.argv[sys.argv.index("--conductor") + 1]
        print(json.dumps({"cadenas": catalogo_cadenas(filtro)},
                         ensure_ascii=False, indent=2))
    elif "--catalogo_conductores" in sys.argv:
        print(json.dumps({"conductores": catalogo_conductores()},
                         ensure_ascii=False, indent=2))
    elif "--sincronizar" in sys.argv:
        idx = sys.argv.index("--sincronizar")
        ruta_peticion = sys.argv[idx + 1]
        ruta_resultado = None
        if "--resultado" in sys.argv:
            ruta_resultado = sys.argv[sys.argv.index("--resultado") + 1]
        with open(ruta_peticion, encoding="utf-8-sig") as archivo:
            p = json.load(archivo)
        try:
            r = sincronizar_tipos_cadenas(
                p.get("excel", ""),
                p.get("apoyos_config"),
                p.get("config_conductor"),
                p.get("salida_apoyos"),
            )
        except Exception as error:
            r = {"ok": False, "error": str(error)}
        if ruta_resultado:
            with open(ruta_resultado, "w", encoding="utf-8") as archivo:
                json.dump(r, archivo, ensure_ascii=False, indent=2, default=str)
        print(json.dumps(r, ensure_ascii=False, indent=2)[:2500])
    elif "--peticion" in sys.argv:
        idx = sys.argv.index("--peticion")
        ruta_peticion = sys.argv[idx + 1]
        ruta_resultado = None
        if "--resultado" in sys.argv:
            ruta_resultado = sys.argv[sys.argv.index("--resultado") + 1]
        r = procesar_peticion_json(ruta_peticion, ruta_resultado)
        print(json.dumps(r, ensure_ascii=False, indent=2)[:2500])
    else:
        print(__doc__)















