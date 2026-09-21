# -*- coding: utf-8 -*-
"""
modulos_fuste.py
================
Motor de edicion del FUSTE de un apoyo: ANADIR o QUITAR MODULOS ENTEROS.

Cada apoyo (Serie C, M50, M60, G, PRESILLA...) se construye con una CABEZA
(armado superior, fija) mas un FUSTE modular de celdas (los "cuadrados"):
  - Serie C          -> cabeza 4.200 m, modulo 0.600 m
  - Serie G          -> cabeza 5.600 m, modulo 0.700 m
  - M50 / M60        -> cabeza 4.000 m, modulo 0.400 m
  - PRESILLA_250     -> cabeza 4.000 m, modulo 0.340 m
  - PRESILLA_400_1250 y trenzados -> cabeza 4.100 m, modulo 0.470 m

Solo se pueden operar MODULOS ENTEROS: 'delta_modulos' es un entero
(positivo anade, negativo quita). Nunca medio modulo.

Efectos de anadir/quitar N modulos de altura 'hm' a un apoyo:
  1. La cabeza (y con ella el punto de amarre de los cables) sube/baja
     delta = N * hm.
  2. Se actualizan TODAS las alturas del Excel:
       - 'Eleccion apoyos'   : Altura de refere. (HREF), Altura libre
                               y Altura libre real.
       - 'Apoyos y crucetas' : Altura normaliz. y Altura total.
       - 'Cimen. monobloque' : Cogolla (altura sobre terreno).
  3. Los DESNIVELES de los dos tramos adyacentes cambian en delta
     (el desnivel es la diferencia de altura entre los puntos de amarre).
  4. Se recalculan TENSIONES y FLECHAS de TODA la linea (EDS constante,
     vanos de regulacion) y los coeficientes L, N, S.
  5. Se escribe en el Excel de trabajo (copia del original que acumula
     cambios) y se actualiza apoyos_configurados.json para que Unity
     reconstruya el apoyo con el numero de modulos correcto.

GENERAL: funciona para cualquier excel de la plantilla. La altura del
modulo se deduce de la referencia del fabricante ('Apoyos y crucetas'),
como hace principal.py.

USO:
    from modulos_fuste import modificar_modulos_fuste
    res = modificar_modulos_fuste("line_example.xlsx", apoyo_numero=2,
                                  delta_modulos=2,
                                  salida_excel="line_example_editada.xlsx")

CLI (para Unity):
    py modulos_fuste.py --peticion peticion_modulos.json --resultado resultado_modulos.json
"""
import math
import os
import re
import json

import openpyxl

from twinelec_paths import assets_dir

import tensiones_vanos as tv

from datos_apoyos_excel import (
    _clave_texto,
    _localizar_columna,
    _numero_decimal,
    _texto_limpio,
)


# ---------------------------------------------------------------------------
# Resolucion de la altura del modulo por familia
# ---------------------------------------------------------------------------
def _resolver_altura_modulo(referencia_apoyo):
    """Familia, altura de cabeza y altura de modulo segun la referencia del
    fabricante. Misma logica que principal.py (_resolver_cuerpo_modular_armado)
    y que BlenderScripts/crear_armados.py (MODULOS_CUERPO)."""
    clave = _clave_texto(referencia_apoyo)
    if "serie c" in clave:
        return "C", 4.200, 0.600
    if "presilla" in clave and "trenz" in clave:
        return "PRESILLA_400_1250_TRENZADOS", 4.100, 0.470
    if "presilla" in clave and re.search(r"\b250\b", clave):
        return "PRESILLA_250", 4.000, 0.340
    if "presilla" in clave:
        return "PRESILLA_400_1250", 4.100, 0.470
    if re.search(r"\bm\s*60\b", clave):
        return "M60", 4.000, 0.400
    if re.search(r"\bm\s*50\b", clave):
        return "M50", 4.000, 0.400
    if re.search(r"\bg\b", clave):
        return "G", 5.600, 0.700
    return None


# ---------------------------------------------------------------------------
# Lectura de las alturas actuales de cada apoyo
# ---------------------------------------------------------------------------
def _fila_de_apoyo(ws, apoyo_numero):
    """Fila (1-based) donde la columna 0 contiene 'apoyo_numero'."""
    for fila in ws.iter_rows(min_row=1, max_col=2):
        numero = _numero_decimal(fila[0].value)
        if numero is not None and int(numero) == int(apoyo_numero):
            return fila[0].row
    return None


def leer_alturas_por_apoyo(ruta_excel):
    """Lee TODAS las alturas de cada apoyo de las tres pestanas:

      - 'Eleccion apoyos'  : Altura de refere. m (HREF), Altura libre m,
                             Altura libre real m.
      - 'Apoyos y crucetas': Referencia del fabricante, Altura normaliz. m,
                             Recrecido cabeza m, Altura total m.
      - 'Cimen. monobloque': Cogolla m (altura sobre terreno).

    Devuelve {apoyo_numero: {campos}}."""
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws_el = wb["Elección apoyos"]
    ws_ay = wb["Apoyos y crucetas"]
    ws_cm = wb["Cimen. monobloque"]

    col_href = _localizar_columna(ws_el, [r"altura.*refere"], fila_cabecera=1)
    col_libre = _localizar_columna(ws_el, [r"^altura libre m"],
                                   fila_cabecera=1)
    col_libre_real = _localizar_columna(ws_el, [r"altura.*libre.*real"],
                                        fila_cabecera=1)
    col_ref = _localizar_columna(ws_ay, [r"referencia.*apoyo"],
                                 fila_cabecera=2)
    col_norm = _localizar_columna(ws_ay, [r"altura.*normali"],
                                  fila_cabecera=2)
    col_rec = _localizar_columna(ws_ay, [r"recrecido"], fila_cabecera=2)
    col_total = _localizar_columna(ws_ay, [r"^altura total"],
                                   fila_cabecera=2)
    col_cogolla = _localizar_columna(ws_cm, [r"cogolla"], fila_cabecera=3)

    resultado = {}
    for fila in ws_el.iter_rows(min_row=1):
        numero = _numero_decimal(fila[0].value)
        if numero is None:
            continue
        apoyo = int(numero)
        resultado.setdefault(apoyo, {})
        if col_href is not None:
            resultado[apoyo]["href_m"] = _numero_decimal(fila[col_href].value)
        if col_libre is not None:
            resultado[apoyo]["altura_libre_m"] = _numero_decimal(
                fila[col_libre].value)
        if col_libre_real is not None:
            resultado[apoyo]["altura_libre_real_m"] = _numero_decimal(
                fila[col_libre_real].value)

    for fila in ws_ay.iter_rows(min_row=1):
        numero = _numero_decimal(fila[0].value)
        if numero is None:
            continue
        apoyo = int(numero)
        resultado.setdefault(apoyo, {})
        if col_ref is not None:
            resultado[apoyo]["referencia_apoyo"] = _texto_limpio(
                fila[col_ref].value)
        if col_norm is not None:
            resultado[apoyo]["altura_normalizada_m"] = _numero_decimal(
                fila[col_norm].value)
        if col_rec is not None:
            resultado[apoyo]["recrecido_cabeza_m"] = _numero_decimal(
                fila[col_rec].value)
        if col_total is not None:
            resultado[apoyo]["altura_total_m"] = _numero_decimal(
                fila[col_total].value)

    for fila in ws_cm.iter_rows(min_row=1):
        numero = _numero_decimal(fila[0].value)
        if numero is None:
            continue
        apoyo = int(numero)
        resultado.setdefault(apoyo, {})
        if col_cogolla is not None:
            resultado[apoyo]["cogolla_m"] = _numero_decimal(
                fila[col_cogolla].value)

    return resultado


def _leer_angulos_excel(ruta_excel):
    """Angulos actuales de la linea desde 'Cálculo de apoyos' (columna
    'Valor ángulo'). Devuelve lista por apoyo en orden 1..N."""
    wb = openpyxl.load_workbook(ruta_excel, data_only=True)
    ws = wb["Cálculo de apoyos"]
    col_ang = _localizar_columna(ws, [r"valor.*angulo"])
    por_numero = {}
    for fila in ws.iter_rows(min_row=1):
        numero = _numero_decimal(fila[0].value)
        if numero is None:
            continue
        v = _numero_decimal(fila[col_ang].value) if col_ang is not None else None
        por_numero[int(numero)] = v if v is not None else 180.0
    if not por_numero:
        return []
    return [por_numero.get(i, 180.0) for i in range(1, max(por_numero) + 1)]


# ---------------------------------------------------------------------------
# Escritura: Excel de trabajo + configuracion de Unity
# ---------------------------------------------------------------------------
def actualizar_alturas_excel(salida_excel, apoyo_numero, alturas):
    """Escribe en el Excel de trabajo las alturas NUEVAS de un apoyo.

    alturas: dict con claves 'href_m', 'altura_libre_m', 'altura_libre_real_m',
    'altura_normalizada_m', 'altura_total_m', 'cogolla_m' (None -> no tocar)."""
    wb = openpyxl.load_workbook(salida_excel)
    ws_el = wb["Elección apoyos"]
    ws_ay = wb["Apoyos y crucetas"]
    ws_cm = wb["Cimen. monobloque"]

    fila_el = _fila_de_apoyo(ws_el, apoyo_numero)
    if fila_el is not None:
        c = {}
        c["href_m"] = _localizar_columna(ws_el, [r"altura.*refere"])
        c["altura_libre_m"] = _localizar_columna(ws_el, [r"^altura libre m"])
        c["altura_libre_real_m"] = _localizar_columna(
            ws_el, [r"altura.*libre.*real"])
        for clave, col in c.items():
            if col is not None and alturas.get(clave) is not None:
                ws_el.cell(row=fila_el, column=col + 1).value = \
                    round(alturas[clave], 2)

    fila_ay = _fila_de_apoyo(ws_ay, apoyo_numero)
    if fila_ay is not None:
        c = {}
        c["altura_normalizada_m"] = _localizar_columna(
            ws_ay, [r"altura.*normali"], fila_cabecera=2)
        c["altura_total_m"] = _localizar_columna(
            ws_ay, [r"^altura total"], fila_cabecera=2)
        for clave, col in c.items():
            if col is not None and alturas.get(clave) is not None:
                ws_ay.cell(row=fila_ay, column=col + 1).value = \
                    round(alturas[clave], 2)

    fila_cm = _fila_de_apoyo(ws_cm, apoyo_numero)
    if fila_cm is not None:
        col = _localizar_columna(ws_cm, [r"cogolla"], fila_cabecera=3)
        if col is not None and alturas.get("cogolla_m") is not None:
            ws_cm.cell(row=fila_cm, column=col + 1).value = \
                round(alturas["cogolla_m"], 2)

    wb.save(salida_excel)
    return salida_excel


def actualizar_config_unity(ruta_config, apoyo_numero, altura_normalizada_m,
                            altura_total_m):
    """Actualiza altura_total_m, altura_normalizada_m y el segmento H#### del
    configuracion_id en apoyos_configurados.json para que Unity reconstruya
    el apoyo con el numero de modulos correcto."""
    with open(ruta_config, encoding="utf-8-sig") as f:
        cfg = json.load(f)
    cambiado = False
    for a in cfg.get("apoyos", []):
        apoyo = a.get("apoyo") or {}
        if apoyo.get("numero") != int(apoyo_numero):
            continue
        apoyo["altura_normalizada_m"] = round(altura_normalizada_m, 3)
        apoyo["altura_total_m"] = round(altura_total_m, 3)
        cid = a.get("configuracion_id", "")
        if cid:
            a["configuracion_id"] = re.sub(
                r"__H\d+$", "__H%04d" % round(altura_total_m * 1000.0), cid)
        cambiado = True
        break
    if not cambiado:
        raise ValueError("No se encontro el apoyo %s en '%s'."
                         % (apoyo_numero, ruta_config))
    with open(ruta_config, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)
    return ruta_config


def _ruta_config_unity(carpeta_excel):
    """Ruta al estado compartido de apoyos que consume Unity."""
    return str(assets_dir() / "apoyos_configurados.json")


# ---------------------------------------------------------------------------
# Orquestador principal
# ---------------------------------------------------------------------------
# ---------------------------------------------------------------------------
# Edicion de modulos
# ---------------------------------------------------------------------------
def _localizar_utm(salida_excel, ruta_excel):
    """Busca el json UTM de la linea: primero junto al Excel de trabajo
    ('utm_editado.json' acumula los cambios ya aplicados), despues junto al
    original y por ultimo en la carpeta del motor (Datos/). NUNCA se modifica
    el archivo de entrada: solo se lee; los cambios van a 'utm_editado.json'."""
    import glob
    base = os.path.splitext(os.path.basename(ruta_excel))[0]
    nombres = ["utm_editado.json", "utm_%s_extraido.json" % base]
    carpetas = [os.path.dirname(os.path.abspath(salida_excel)),
                os.path.dirname(os.path.abspath(ruta_excel)),
                os.path.dirname(os.path.abspath(__file__))]
    for carpeta in carpetas:
        for nombre in nombres:
            ruta = os.path.join(carpeta, nombre)
            if os.path.exists(ruta):
                return ruta
        globs = glob.glob(os.path.join(carpeta, "utm_*.json"))
        if globs:
            return globs[0]
    return None


def modificar_modulos_fuste(ruta_excel, apoyo_numero, delta_modulos,
                            altura_modulo_m=None, salida_excel=None,
                            viento_kmh=120.0, config_unity=None,
                            min_modulos=1, ruta_utm=None):
    """Anade (delta>0) o quita (delta<0) MODULOS ENTEROS del fuste de un
    apoyo y recalcula toda la linea (desniveles, tensiones, flechas, L/N/S)
    escribiendo el resultado en el Excel de trabajo y en apoyos_configurados.json.

    Reglas:
      - delta_modulos: entero distinto de 0 (nunca medio modulo).
      - altura_modulo_m: si es None se deduce de la referencia del fabricante.
      - min_modulos: numero minimo de celdas de fuste que deben quedar (>=1).
      - ruta_excel: el original o el '_editada.xlsx' que acumula cambios
        (recomendado pasar el editada si existe, como hace la ventana de Unity).

    Devuelve un dict resumen con las alturas y la geometria recalculada.
    """
    import datos_apoyos_excel as dae
    from sobrecargas import calcular_sobrecargas
    import mover_apoyo as ma

    # ---- Validaciones de "modulo entero" ----
    if isinstance(delta_modulos, bool) or delta_modulos is None:
        raise ValueError("delta_modulos debe ser un entero distinto de 0.")
    try:
        n_mod = float(delta_modulos)
    except (TypeError, ValueError):
        raise ValueError("delta_modulos debe ser un numero entero. Recibido: %r"
                         % (delta_modulos,))
    if abs(n_mod - round(n_mod)) > 1e-9 or round(n_mod) == 0:
        raise ValueError(
            "Solo se pueden anadir o quitar MODULOS ENTEROS (1, 2, 3...). "
            "Recibido: %r." % (delta_modulos,))
    n_mod = int(round(n_mod))

    if salida_excel is None:
        base = os.path.splitext(os.path.basename(ruta_excel))[0]
        salida_excel = os.path.join(os.path.dirname(os.path.abspath(ruta_excel)),
                                    base + "_editada.xlsx")

    # RECIPROCIDAD: si ya existe el Excel de trabajo (acumula movimientos y
    # modulos), se lee el estado ACTUAL de ahi, no del original.
    salida_abs = os.path.abspath(salida_excel)
    ruta_abs = os.path.abspath(ruta_excel)
    fuente = salida_abs if (os.path.exists(salida_abs)
                            and salida_abs != ruta_abs) else ruta_abs

    # ---- Alturas actuales del apoyo ----
    alturas = leer_alturas_por_apoyo(fuente)
    a = alturas.get(int(apoyo_numero))
    if a is None:
        raise ValueError("No se encontraron alturas para el apoyo %s en '%s'."
                         % (apoyo_numero, ruta_excel))

    if altura_modulo_m is None:
        resolucion = _resolver_altura_modulo(a.get("referencia_apoyo"))
        if resolucion is None:
            raise ValueError(
                "No se pudo deducir la altura del modulo del apoyo %s "
                "(referencia: %s). Pasa 'altura_modulo_m'."
                % (apoyo_numero, a.get("referencia_apoyo")))
        familia, altura_cabeza, altura_modulo_m = resolucion
    else:
        familia, altura_cabeza = None, None

    altura_modulo_m = float(altura_modulo_m)
    if altura_modulo_m <= 0:
        raise ValueError("altura_modulo_m debe ser > 0.")

    delta_altura = n_mod * altura_modulo_m

    # Base de alturas: si alguna falta se sustituye por otra presente.
    base_total = (a.get("cogolla_m") or a.get("altura_total_m")
                  or a.get("href_m"))
    base_href = (a.get("href_m") or a.get("altura_total_m")
                 or a.get("cogolla_m"))
    if base_total is None or base_href is None:
        raise ValueError("No se pudieron leer las alturas del apoyo %s."
                         % apoyo_numero)

    nueva_cogolla = base_total + delta_altura
    nueva_href = base_href + delta_altura
    nueva_norm = ((a.get("altura_normalizada_m") or nueva_href)
                  + delta_altura)
    nueva_total_ac = ((a.get("altura_total_m") or nueva_cogolla)
                      + delta_altura)
    nueva_libre = ((a.get("altura_libre_m") + delta_altura)
                   if a.get("altura_libre_m") is not None else None)
    nueva_libre_real = ((a.get("altura_libre_real_m") + delta_altura)
                        if a.get("altura_libre_real_m") is not None else None)

    # Restriccion: no se puede quitar mas modulos de los que hay.
    if altura_cabeza is None:
        resolucion = _resolver_altura_modulo(a.get("referencia_apoyo"))
        altura_cabeza = resolucion[1] if resolucion else 4.0
    altura_cabeza = float(altura_cabeza)
    minimo_total = altura_cabeza + min_modulos * altura_modulo_m
    if nueva_cogolla < minimo_total - 1e-6 or nueva_href < minimo_total - 1e-6:
        max_quitar = int(math.floor((base_total - minimo_total)
                                    / altura_modulo_m))
        raise ValueError(
            "No se pueden %s %d modulo(s): el apoyo %d quedaria con %.2f m de "
            "altura total, por debajo del minimo de %.2f m (cabeza %.2f m + "
            "al menos %d modulo de %.2f m). Maximo extraible: %d modulo(s)."
            % ("quitar" if n_mod < 0 else "anadir", abs(n_mod), apoyo_numero,
               nueva_cogolla, minimo_total, altura_cabeza, min_modulos,
               altura_modulo_m, max_quitar))
    # ---- Geometria: solo cambian los desniveles de los tramos aledanos ----
    geom = ma.leer_geometria_excel(fuente)
    vanos = [t[0] for t in geom["tramos"]]
    desniveles = [t[1] for t in geom["tramos"]]
    k = int(apoyo_numero) - 1
    if k < 0 or k >= len(vanos) + 1:
        raise ValueError("apoyo_numero fuera de rango (1..%d)"
                         % (len(vanos) + 1))
    if k - 1 >= 0:
        desniveles[k - 1] += delta_altura
    if k < len(vanos):
        desniveles[k] -= delta_altura

    angulos = _leer_angulos_excel(fuente)
    if len(angulos) != len(vanos) + 1:
        angulos = [180.0] * (len(vanos) + 1)
    coefs = ma.coeficientes_apoyos(vanos, desniveles, angulos)

    # ---- Conductor y sobrecargas por tramo ----
    conductor = dae.leer_conductor(fuente)
    props = tv.propiedades_conductor(conductor)
    alfa = props["alfa_1_C"]
    E = props["E_daN_mm2"]
    S = props["seccion_mm2"]
    p0 = props["peso_daN_m"]
    peso_kg_m = conductor["masa_kg_km"] / 1000.0
    diam_mm = conductor["diametro_mm"]

    zonas_tramo = dae.extraer_zonas_por_tramo(fuente)
    sobrecarga_por_tramo = []
    for i in range(len(vanos)):
        zona = zonas_tramo.get((i + 1, i + 2), "A")
        sob = calcular_sobrecargas(diam_mm, peso_kg_m, zona, viento_kmh)
        sobrecarga_por_tramo.append({
            "Sv": sob["Sv"], "Sv120": sob["Sv120"], "Sh": sob["Sh"],
            "S_vm": sob["S_vm"], "Peso cable": p0,
        })

    # ---- EDS (opcion A: constante) y tensiones por vano de regulacion ----
    estado_eds = ma.leer_estado_eds(fuente)
    estado_eds["p0_daN_m"] = p0
    estado_eds["T0_por_tramo"] = [
        estado_eds["T0_por_tramo"].get((i + 1, i + 2))
        for i in range(len(vanos))]

    grupos_clave = ma.leer_grupos_vanos_regulacion(fuente)
    grupos_por_indice = {}
    grupos_tramos = {}
    for i in range(len(vanos)):
        g = grupos_clave.get((i + 1, i + 2), i + 1)
        grupos_por_indice[i] = g
        grupos_tramos.setdefault(g, []).append(i)
    vano_reg_por_tramo = []
    for i in range(len(vanos)):
        ar_grupo = tv.vano_regulacion(
            [vanos[j] for j in grupos_tramos[grupos_por_indice[i]]])
        vano_reg_por_tramo.append(ar_grupo if ar_grupo is not None else vanos[i])

    tensiones = ma.recalcular_tensiones_tramos(
        vanos, desniveles, estado_eds, alfa, E, S, sobrecarga_por_tramo,
        grupos=grupos_por_indice)

    # ---- Excel de trabajo (copia acumulativa) ----
    ma.copiar_excel_trabajo(ruta_excel, salida_excel)
    actualizar_alturas_excel(salida_excel, int(apoyo_numero), {
        "href_m": nueva_href,
        "altura_libre_m": nueva_libre,
        "altura_libre_real_m": nueva_libre_real,
        "altura_normalizada_m": nueva_norm,
        "altura_total_m": nueva_total_ac,
        "cogolla_m": nueva_cogolla,
    })
    ma.actualizar_excel_trabajo(salida_excel, vanos, desniveles, angulos,
                                coefs, tensiones,
                                vano_reg_por_tramo=vano_reg_por_tramo)

    # ---- Esfuerzos V/T/L de las hipotesis (CalculoApoyos) ----
    try:
        from cambiar_cadena_conductor import recalcular_esfuerzos_apoyos
        recalcular_esfuerzos_apoyos(ruta_excel, salida_excel,
                                    tensiones=tensiones)
        print("  [modulos_fuste] Esfuerzos V/T/L actualizados.")
    except Exception as error_esf:
        print("  ⚠️ No se pudieron recalcular los esfuerzos V/T/L:",
              error_esf)

    # ---- Configuracion de Unity (reconstruye el apoyo) ----
    ruta_config = None
    if config_unity is None:
        config_unity = _ruta_config_unity(
            os.path.dirname(os.path.abspath(ruta_excel)))
    if os.path.exists(config_unity):
        ruta_config = actualizar_config_unity(
            config_unity, int(apoyo_numero), nueva_norm, nueva_cogolla)
        print("  [modulos_fuste] Unity actualizado:", ruta_config)

    # ---- COTA UTM del apoyo: cambia en el mismo delta de altura. La cota
    #      sobre el nivel del mar del apoyo sube/baja con el modulo, y con
    #      ella los desniveles (ya recalculados arriba) y las tensiones.
    #      Se escribe en 'utm_editado.json' (junto al Excel de trabajo);
    #      los archivos de ENTRADA no se tocan. ----
    # IMPORTANTE: se PREFIERE leer de 'utm_editado.json' si ya existe (acumula
    # los cambios aplicados en pulsaciones anteriores). La ventana de Unity
    # cachea la ruta UTM al abrirse y puede apuntar al json ORIGINAL, con lo
    # que sin esta preferencia cada +1MOD releeria la cota inicial y no
    # acumularia (+0,6, +0,6, +0,6...). ----
    aviso_utm = None
    cota_utm_antes = cota_utm_despues = None
    ruta_utm_editado = None
    try:
        utm_editado_ruta = os.path.join(
            os.path.dirname(os.path.abspath(salida_excel)),
            "utm_editado.json")
        ruta_lectura = utm_editado_ruta if os.path.exists(utm_editado_ruta) \
            else (ruta_utm if ruta_utm and os.path.exists(ruta_utm) else None)
        if ruta_lectura is None:
            ruta_lectura = _localizar_utm(salida_excel, ruta_excel)
        if ruta_lectura and os.path.exists(ruta_lectura):
            utm_actual = ma.leer_utm(ruta_lectura)
            kk = int(apoyo_numero) - 1
            if 0 <= kk < len(utm_actual):
                cota_utm_antes = round(utm_actual[kk]["cota"], 3)
                utm_actual[kk]["cota"] = (float(utm_actual[kk]["cota"])
                                          + delta_altura)
                cota_utm_despues = round(utm_actual[kk]["cota"], 3)
                with open(utm_editado_ruta, "w", encoding="utf-8") as f:
                    json.dump(utm_actual, f, ensure_ascii=False, indent=2)
                ma.regenerar_csv_linea(utm_editado_ruta)
                try:
                    from generar_doc_utm import generar_docx_desde_json
                    generar_docx_desde_json(utm_editado_ruta)
                except Exception as error_doc:
                    print("  ⚠️ Documento UTM no regenerado:", error_doc)
                ruta_utm_editado = utm_editado_ruta
            else:
                aviso_utm = "apoyo fuera de rango en el json UTM"
        else:
            aviso_utm = "no se encontro un json UTM para la linea"
    except Exception as error_utm:
        aviso_utm = str(error_utm)

    return {
        "apoyo_numero": int(apoyo_numero),
        "familia": familia,
        "delta_modulos": n_mod,
        "altura_modulo_m": altura_modulo_m,
        "delta_altura_m": round(delta_altura, 3),
        "altura_total_antes": round(base_total, 3),
        "altura_total_despues": round(nueva_cogolla, 3),
        "href_antes": round(base_href, 3),
        "href_despues": round(nueva_href, 3),
        "modulos_fuste": _numero_modulos_visibles(
            nueva_cogolla, altura_cabeza, altura_modulo_m),
        "cota_utm_antes": cota_utm_antes,
        "cota_utm_despues": cota_utm_despues,
        "aviso_utm": aviso_utm,
        "utm_editado": ruta_utm_editado,
        "salida_excel": os.path.abspath(salida_excel),
        "config_unity": os.path.abspath(ruta_config) if ruta_config else None,
        "vanos": vanos,
        "desniveles": desniveles,
        "angulos": angulos,
        "coeficientes": coefs,
        "tensiones_por_tramo": tensiones,
    }


def _numero_modulos_visibles(altura_total_m, altura_cabeza_m, altura_modulo_m):
    """Numero de celdas (cuadrados) visibles que dibuja Unity:
    ceil((altura_total - cabeza) / modulo)."""
    cuerpo = max(0.0, altura_total_m - altura_cabeza_m)
    if altura_modulo_m <= 0:
        return 0
    return int(math.ceil(cuerpo / altura_modulo_m))


# ---------------------------------------------------------------------------
# CLI (JSON) para Unity
# ---------------------------------------------------------------------------
def procesar_peticion_json(ruta_peticion, ruta_resultado=None):
    """Lee la peticion JSON de Unity, anade/quita modulos y escribe el
    resultado (o el error). Peticion:
        {"excel":.., "utm":.., "salida_excel":.., "apoyo":N,
         "modulos":+2|-1, "altura_modulo_m":null, "config_unity":..}
    'modulos' es un entero: positivo anade, negativo quita."""
    with open(ruta_peticion, encoding="utf-8-sig") as f:
        p = json.load(f)
    try:
        res = modificar_modulos_fuste(
            p.get("excel"), int(p["apoyo"]), p.get("modulos"),
            altura_modulo_m=(p.get("altura_modulo_m")
                             if p.get("altura_modulo_m") is not None
                             else None),
            salida_excel=p.get("salida_excel"),
            config_unity=p.get("config_unity"),
            ruta_utm=p.get("utm"))
        resultado = {"ok": True}
        resultado.update(res)
    except Exception as e:
        resultado = {"ok": False, "error": str(e)}
    if ruta_resultado:
        with open(ruta_resultado, "w", encoding="utf-8") as f:
            json.dump(resultado, f, ensure_ascii=False, indent=2, default=str)
    return resultado


if __name__ == "__main__":
    import sys
    import io
    try:
        if sys.stdout.encoding and "utf" not in sys.stdout.encoding.lower():
            sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                                          errors="replace")
    except Exception:
        pass

    if "--peticion" in sys.argv:
        idx = sys.argv.index("--peticion")
        ruta_peticion = sys.argv[idx + 1]
        ruta_resultado = None
        if "--resultado" in sys.argv:
            ruta_resultado = sys.argv[sys.argv.index("--resultado") + 1]
        r = procesar_peticion_json(ruta_peticion, ruta_resultado)
        print(json.dumps(r, ensure_ascii=False, indent=2)[:2500])
    else:
        CARPETA = os.path.dirname(os.path.abspath(__file__))
        RUTA = os.path.join(CARPETA, "line_example.xlsx")
        SALIDA = os.path.join(CARPETA, "line_example_editada.xlsx")

        print("== DEMO: anadir 2 modulos de fuste al apoyo 2 ==")
        res = modificar_modulos_fuste(RUTA, apoyo_numero=2, delta_modulos=2,
                                      salida_excel=SALIDA)
        print("  Excel de trabajo:       ", res["salida_excel"])
        print("  Familia: %s | modulo %.3f m | delta %.3f m"
              % (res["familia"], res["altura_modulo_m"], res["delta_altura_m"]))
        print("  Altura total: %.2f m -> %.2f m (modulos fuste: %d)"
              % (res["altura_total_antes"], res["altura_total_despues"],
                 res["modulos_fuste"]))
        print("  HREF: %.2f m -> %.2f m"
              % (res["href_antes"], res["href_despues"]))
        print("  Desniveles nuevos:      ", [round(d, 3)
                                             for d in res["desniveles"]])
        print("  Tensiones (T en daN):")
        for i, t in res["tensiones_por_tramo"].items():
            fila = " ".join("%s=%s" % (k, round(v["T"], 1))
                            for k, v in t.items())
            print("    tramo %d-%d: %s" % (i + 1, i + 2, fila))




